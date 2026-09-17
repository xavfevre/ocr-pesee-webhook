# -*- coding: utf-8 -*-
"""Fichier de travail « Tarifs Maquignon » : articles avec prix actuels par liste (public, pro, pierres), règles clients,
clients et liste actuelle / proposée, listes existantes et action proposée. Lecture seule sur Odoo."""
import os, ssl, sys, xmlrpc.client, collections, datetime
sys.stdout.reconfigure(encoding='utf-8')
U = 'https://maquignon.odoo.com'; D = 'maquignon'; us = os.environ['ODOO_USER']; p = os.environ['ODOO_PWD']
c = ssl.create_default_context()
uid = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/common', context=c).authenticate(D, us, p, {})
m = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/object', context=c)
x = lambda mo, me, *a, **k: m.execute_kw(D, uid, p, mo, me, list(a), k)
ctx = {'allowed_company_ids': [1], 'active_test': False}
today = datetime.date.today()
PUBLIC, PRO, PIERRES, ANC_PRO, PARTICULIER = 302, 299, 1673, 12, 1
# ---- listes
pls = x('product.pricelist', 'search_read', [['company_id', 'in', [1, False]]], fields=['name', 'active', 'item_ids', 'sequence'], order='sequence,id', context=ctx)
plname = {pl['id']: pl['name'] for pl in pls}
items = x('product.pricelist.item', 'search_read', [['pricelist_id', 'in', list(plname)]],
          fields=['pricelist_id', 'applied_on', 'product_tmpl_id', 'product_id', 'categ_id', 'compute_price', 'fixed_price', 'percent_price', 'base',
                  'base_pricelist_id', 'price_discount', 'price_surcharge', 'price_round', 'min_quantity', 'date_start', 'date_end'], limit=20000, context=ctx)
byp = collections.defaultdict(list)
for it in items:
    byp[it['pricelist_id'][0]].append(it)
# ---- clients (fiches principales) : clients ou rattachés à une liste
civ = x('ir.model.fields', 'search_read', [['model', '=', 'res.partner'], ['field_description', 'ilike', 'Civilit'], ['ttype', 'in', ['many2one', 'selection', 'char']]], fields=['name', 'ttype'])
CIV = civ[0]['name'] if civ else None
print('champ civilité :', civ)
partners = x('res.partner', 'search_read', ['|', ['customer_rank', '>', 0], ['specific_property_product_pricelist', '!=', False], ['parent_id', '=', False]],
             fields=['name', 'is_company', 'vat', 'specific_property_product_pricelist', 'category_id', 'sale_order_count', 'customer_rank', 'city', 'ref', 'active'] + ([CIV] if CIV else []), limit=10000, context=ctx)
partners = [pt for pt in partners if pt['active']]
part_by_pl = collections.defaultdict(list)
for pt in partners:
    if pt['specific_property_product_pricelist']:
        part_by_pl[pt['specific_property_product_pricelist'][0]].append(pt['name'])
# ---- articles vendables : variantes
prods = x('product.product', 'search_read', [['sale_ok', '=', True], ['company_id', 'in', [1, False]], ['active', '=', True]],
          fields=['default_code', 'display_name', 'lst_price', 'list_price', 'standard_price', 'categ_id', 'uom_id', 'type', 'product_tmpl_id'], limit=5000, context=ctx)
prods.sort(key=lambda pr: ((pr['categ_id'] and pr['categ_id'][1]) or '', pr['display_name']))
pids = [pr['id'] for pr in prods]
print('articles (variantes) vendables :', len(prods))


cats = {ct['id']: ct['parent_path'] for ct in x('product.category', 'search_read', [], fields=['parent_path'], context=ctx)}
ORDER = {'0_product_variant': 0, '1_product': 1, '2_product_category': 2, '3_global': 3}
today_s = today.isoformat()


def regles(pl_id):
    its = [it for it in byp[pl_id] if (not it['date_start'] or it['date_start'][:10] <= today_s) and (not it['date_end'] or it['date_end'][:10] >= today_s) and (it['min_quantity'] or 0) <= 1]
    return sorted(its, key=lambda it: (ORDER[it['applied_on']], -(it['min_quantity'] or 0), -(it['categ_id'][0] if it['categ_id'] else 0), -it['id']))


def prix(pl_id, pr, depth=0):
    """Reproduit product.pricelist._compute_price_rule pour 1 unité (règles fixe / % / formule, base autre liste)."""
    if depth > 5:
        return pr['lst_price']
    for it in regles(pl_id):
        a = it['applied_on']
        if a == '0_product_variant' and (not it['product_id'] or it['product_id'][0] != pr['id']):
            continue
        if a == '1_product' and (not it['product_tmpl_id'] or it['product_tmpl_id'][0] != pr['product_tmpl_id'][0]):
            continue
        if a == '2_product_category' and not (it['categ_id'] and pr['categ_id'] and cats.get(pr['categ_id'][0], '').startswith(cats.get(it['categ_id'][0], '#'))):
            continue
        base = it.get('base') or 'list_price'
        if base == 'pricelist' and it.get('base_pricelist_id'):
            bp = prix(it['base_pricelist_id'][0], pr, depth + 1)
        elif base == 'standard_price':
            bp = pr['standard_price']
        else:
            bp = pr['lst_price']
        cp = it['compute_price']
        if cp == 'fixed':
            return it['fixed_price'] or 0.0
        if cp == 'percentage':
            return bp * (1 - (it['percent_price'] or 0) / 100)
        v = bp * (1 - (it.get('price_discount') or 0) / 100)
        if it.get('price_round'):
            v = round(v / it['price_round']) * it['price_round']
        return v + (it.get('price_surcharge') or 0)
    return pr['lst_price']


def prix_liste(pl_id):
    return {pr['id']: round(prix(pl_id, pr), 2) for pr in prods}


px = {pl: prix_liste(pl) for pl in (PUBLIC, PRO, PIERRES)}
# ventes 12 mois par variante (quantité, montant)
depuis = (today - datetime.timedelta(days=365)).isoformat()
ventes = {}
for r in x('sale.order.line', 'read_group', [['order_id.company_id', '=', 1], ['order_id.date_order', '>=', depuis], ['state', '=', 'sale'], ['display_type', '=', False], ['product_id', 'in', pids]],
                          ['product_uom_qty:sum', 'price_subtotal:sum'], ['product_id'], lazy=False, context=ctx):
    ventes[r['product_id'][0]] = (r['product_uom_qty'], r['price_subtotal'], r['__count'])
# ---- classeur
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
wb = Workbook()
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='01666B'); SAISIE = PatternFill('solid', fgColor='FFF2CC')


def entete(ws, cols, widths):
    ws.append(cols)
    for cell in ws[1]:
        cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 32
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'


ws = wb.active; ws.title = 'Lisez-moi'
ws.column_dimensions['A'].width = 120
for line in [
    'TARIFS SARL MAQUIGNON — fichier de travail du %s' % today.strftime('%d/%m/%Y'),
    '',
    'Situation : 30 listes de prix, dont 2 archivées encore utilisées (« Anciens Tarifs Pro » : 224 clients sans aucune règle, « Prix Particulier » : 13 clients),',
    '17 listes « client » (BESNAULT, TERROBA, GSM, LEFEVRE…, COLAS…) dont plusieurs se basent sur « Anciens Tarifs Pro » (vide), et 5 listes sans client.',
    '',
    'Structure proposée :',
    '  1. « Tarif Particulier » = prix public de la fiche article (aucune règle, ou une règle 0 % sur le prix public).',
    '  2. « Tarif Professionnel » = un prix fixe par article (repris de « Tarif Pro 2026 » puis corrigé ici), ou une remise en % par catégorie.',
    '  3. Prix spécifiques client = une petite liste au nom du client : ses quelques prix fixes + une règle « tout le reste : Tarif Professionnel ».',
    '     Chaque client n\'a qu\'une seule liste dans Odoo, c\'est cette règle de repli qui lui donne le tarif pro sur tout le reste.',
    '  4. Les autres listes sont archivées ; les clients sont rattachés à Particulier, Professionnel ou leur liste client.',
    '',
    'Comment remplir :',
    '  - Onglet « Articles » : colonnes jaunes « Particulier » et « Professionnel » à vérifier ou corriger (pré-remplies avec les prix actuels). Vide = pas de prix pro spécifique (le pro paie le prix particulier).',
    '  - Onglet « Prix clients » : règles actuelles des listes client. Colonne « Garder » : O pour garder, N pour supprimer. Ajouter des lignes pour de nouveaux prix client (Client, Article, Prix).',
    '  - Onglet « Clients » : colonne « Tarif proposé » à corriger si besoin (Particulier / Professionnel / nom du client pour un tarif spécifique).',
    '  - Onglet « Listes actuelles » : pour information, avec l\'action proposée.',
    'Renvoie-moi le fichier rempli : je crée les listes et les règles dans Odoo, je rattache les clients et j\'archive le reste (réversible).',
]:
    ws.append([line])
ws['A1'].font = Font(bold=True, size=13)
# Articles
wa = wb.create_sheet('Articles')
entete(wa, ['Réf.', 'Article', 'Catégorie', 'Unité', 'Prix public fiche', 'Tarif Public 2026 (actuel)', 'Tarif Pro 2026 (actuel)', 'Prix Public - Pierres (actuel)',
            'Qté vendue 12 mois', 'CA 12 mois HT', 'Nb lignes 12 mois', 'PARTICULIER (à saisir)', 'PROFESSIONNEL (à saisir)', 'Remise pro % (calcul)'],
       (12, 52, 34, 9, 12, 14, 14, 14, 11, 12, 10, 14, 16, 12))
for pr in prods:
    v = ventes.get(pr['id'], (0, 0, 0))
    pu = px[PUBLIC].get(pr['id'], pr['lst_price']); pp = px[PRO].get(pr['id']); ppi = px[PIERRES].get(pr['id'])
    row = [pr['default_code'] or '', pr['display_name'], (pr['categ_id'] and pr['categ_id'][1]) or '', pr['uom_id'][1], pr['lst_price'], pu, pp, ppi,
           round(v[0], 3), round(v[1], 2), v[2], pu, pp if (pp is not None and abs(pp - pu) > 0.005) else None, None]
    wa.append(row)
    r = wa.max_row
    wa.cell(r, 12).fill = SAISIE; wa.cell(r, 13).fill = SAISIE
    wa.cell(r, 14).value = '=IF(AND(L%d>0,M%d<>""),ROUND((L%d-M%d)/L%d*100,1),"")' % (r, r, r, r, r)
wa.auto_filter.ref = wa.dimensions
# Prix clients
wc = wb.create_sheet('Prix clients')
entete(wc, ['Liste', 'Clients rattachés', 'Article', 'S\'applique à', 'Règle', 'Prix fixe', 'Remise %', 'Base', 'Qté min', 'Début', 'Fin', 'GARDER (O/N)', 'NOUVEAU PRIX (si changement)'],
       (30, 40, 50, 14, 12, 10, 9, 22, 8, 11, 11, 12, 16))
for pl in pls:
    if pl['id'] in (PUBLIC, PRO, PIERRES, ANC_PRO, PARTICULIER) or not byp[pl['id']]:
        continue
    for it in byp[pl['id']]:
        cible = (it.get('product_id') or it.get('product_tmpl_id') or it.get('categ_id') or [0, 'Tous les articles'])[1]
        base = (it.get('base_pricelist_id') and it['base_pricelist_id'][1]) or {'list_price': 'prix public', 'standard_price': 'coût', 'pricelist': 'autre liste'}.get(it.get('base'), it.get('base'))
        wc.append([pl['name'], ', '.join(part_by_pl.get(pl['id'], [])) or '(aucun)', cible, {'0_product_variant': 'variante', '1_product': 'article', '2_product_category': 'catégorie', '3_global': 'tout'}.get(it['applied_on'], it['applied_on']),
                   it['compute_price'], it.get('fixed_price') if it['compute_price'] == 'fixed' else None,
                   it.get('percent_price') if it['compute_price'] == 'percentage' else (it.get('price_discount') if it['compute_price'] == 'formula' else None),
                   base if it['compute_price'] == 'formula' else None, it.get('min_quantity'), it.get('date_start') or '', it.get('date_end') or '', 'O', None])
        r = wc.max_row; wc.cell(r, 12).fill = SAISIE; wc.cell(r, 13).fill = SAISIE
for _ in range(15):
    wc.append(['(nouveau) client :', '', '(article)', 'article', 'fixed', None, None, None, 0, '', '', 'O', None])
    r = wc.max_row
    for col in (1, 3, 6):
        wc.cell(r, col).fill = SAISIE
wc.auto_filter.ref = 'A1:M%d' % wc.max_row
# Clients
wk = wb.create_sheet('Clients')
entete(wk, ['Client', 'Réf.', 'Société ?', 'Civilité', 'TVA', 'Ville', 'Étiquettes', 'Nb commandes', 'Liste actuelle', 'TARIF PROPOSÉ (à corriger)'], (36, 12, 9, 10, 16, 18, 24, 10, 34, 30))
PRO_LIKE = {PRO, ANC_PRO, 300, 301, 303}
for pt in sorted(partners, key=lambda pt: (-(pt['sale_order_count'] or 0), pt['name'])):
    pl = pt['specific_property_product_pricelist']; plid = pl and pl[0]
    if plid in PRO_LIKE:
        prop = 'Professionnel'
    elif plid in (PUBLIC, PARTICULIER, None, False):
        prop = 'Professionnel' if (pt['is_company'] and pt['vat']) else 'Particulier'
    else:
        prop = 'Client : ' + plname.get(plid, pl[1])
    wk.append([pt['name'], pt['ref'] or '', 'Oui' if pt['is_company'] else 'Non', ((pt.get(CIV) and (pt[CIV][1] if isinstance(pt[CIV], list) else pt[CIV])) or '') if CIV else '', pt['vat'] or '', pt['city'] or '',
               ', '.join(x('res.partner.category', 'read', pt['category_id'], fields=['name'], context=ctx)[0]['name'] for _ in [0]) if False else '', pt['sale_order_count'], (pl and pl[1]) or '(aucune : Tarif Public 2026)', prop])
    wk.cell(wk.max_row, 10).fill = SAISIE
wk.auto_filter.ref = wk.dimensions
# étiquettes des clients (une seule lecture)
cats = {ct['id']: ct['name'] for ct in x('res.partner.category', 'search_read', [], fields=['name'], context=ctx)}
for i, pt in enumerate(sorted(partners, key=lambda pt: (-(pt['sale_order_count'] or 0), pt['name'])), start=2):
    wk.cell(i, 7).value = ', '.join(cats.get(cid, '') for cid in pt['category_id'])
# Listes actuelles
wl = wb.create_sheet('Listes actuelles')
entete(wl, ['Id', 'Liste', 'Active', 'Règles', 'Clients rattachés', 'Détail des règles', 'Action proposée'], (6, 40, 8, 8, 10, 60, 50))
for pl in pls:
    its = byp[pl['id']]
    det = collections.Counter('%s / %s%s' % (it['applied_on'][2:], it['compute_price'], (' → ' + it['base_pricelist_id'][1]) if it.get('base_pricelist_id') else '') for it in its)
    n = len(part_by_pl.get(pl['id'], []))
    if pl['id'] == PUBLIC:
        act = 'Devient « Tarif Particulier » (renommer)'
    elif pl['id'] == PRO:
        act = 'Devient « Tarif Professionnel » (renommer, prix corrigés depuis l\'onglet Articles)'
    elif pl['id'] in (ANC_PRO, PARTICULIER):
        act = 'Archivée : clients à rattacher (%d)' % n
    elif pl['id'] == PIERRES:
        act = 'Grille de référence pierres : à fondre dans le prix public des fiches, puis archiver'
    elif n == 0:
        act = 'Aucun client : archiver'
    else:
        act = 'Liste client : garder ses prix fixes + repli sur Tarif Professionnel (onglet Prix clients)'
    wl.append([pl['id'], pl['name'], 'Oui' if pl['active'] else 'Non', len(its), n, '; '.join('%s ×%d' % (k, v) for k, v in det.items()), act])
out = 'C:/Users/xavfe/Desktop/Maquignon/Tarifs_Maquignon_%s.xlsx' % today.strftime('%Y-%m-%d')
wb.save(out)
print('fichier :', out)
# stats utiles
diff_pub = sum(1 for pr in prods if px[PUBLIC].get(pr['id']) is not None and abs(px[PUBLIC][pr['id']] - pr['lst_price']) > 0.005)
pro_diff = sum(1 for pr in prods if px[PRO].get(pr['id']) is not None and abs(px[PRO][pr['id']] - px[PUBLIC].get(pr['id'], pr['lst_price'])) > 0.005)
pierres_diff = sum(1 for pr in prods if px[PIERRES].get(pr['id']) is not None and abs(px[PIERRES][pr['id']] - pr['lst_price']) > 0.005)
print('Tarif Public 2026 ≠ prix fiche :', diff_pub, '| Tarif Pro ≠ public :', pro_diff, '| Prix Public - Pierres ≠ prix fiche :', pierres_diff, '| prix fiche à 0 :', sum(1 for pr in prods if not pr['lst_price']))
print('clients dans l onglet :', len(partners), '| par proposition :', dict(collections.Counter(wk.cell(i, 10).value.split(' : ')[0] for i in range(2, wk.max_row + 1))))
print('ventes 12 mois : variantes vendues', len(ventes), 'sur', len(prods))

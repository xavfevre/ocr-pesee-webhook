# -*- coding: utf-8 -*-
"""Lecture seule : bons de commande SARL MAQUIGNON confirmés, jamais facturés, composés uniquement d'articles de
transport / location / transfert (les mêmes que le ménage des BL). Synthèse par mois, famille, client, créateur et
tâche de transport liée ; Excel détaillé sur le Bureau."""
import os, ssl, sys, xmlrpc.client, collections, datetime
sys.stdout.reconfigure(encoding='utf-8')
U = 'https://maquignon.odoo.com'; D = 'maquignon'; us = os.environ['ODOO_USER']; p = os.environ['ODOO_PWD']
c = ssl.create_default_context()
uid = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/common', context=c).authenticate(D, us, p, {})
m = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/object', context=c)
x = lambda mo, me, *a, **k: m.execute_kw(D, uid, p, mo, me, list(a), k)
CTX = {'allowed_company_ids': [1, 2, 3, 4, 13], 'active_test': False}
MOTS = ['transport', 'location de mat', 'transfert de mat', 'livraison']
prods = x('product.product', 'search_read', [['type', '=', 'consu'], ['is_storable', '=', False], '|', '|', '|'] + [['name', 'ilike', w] for w in MOTS],
          fields=['display_name'], context=CTX)
pids = {pr['id'] for pr in prods}


def famille(nom):
    n = (nom or '').lower()
    for k in ('transport de granulats', 'transport appro', 'transport de pierres', 'transport de blocs', 'location de mat', 'transfert de mat'):
        if k in n:
            return k.capitalize()
    return 'Autre transport / livraison'


AVANT = sys.argv[1] if len(sys.argv) > 1 else None   # ex. 2026-09-01 : commandes datées avant cette date seulement
orders = x('sale.order', 'search_read', [['company_id', '=', 1], ['state', '=', 'sale'], ['invoice_ids', '=', False]] + ([['date_order', '<', AVANT]] if AVANT else []),
           fields=['name', 'date_order', 'partner_id', 'amount_untaxed', 'invoice_status', 'origin', 'client_order_ref', 'user_id', 'create_uid', 'create_date',
                   'x_studio_tache', 'x_studio_transport_adresse', 'x_studio_transport_commentaire', 'x_studio_a_facturer', 'picking_ids', 'order_line', 'note'],
           limit=10000, context=CTX)
oids = [o['id'] for o in orders]
lines = x('sale.order.line', 'search_read', [['order_id', 'in', oids], ['display_type', '=', False]],
          fields=['order_id', 'product_id', 'name', 'product_uom_qty', 'qty_delivered', 'qty_invoiced', 'qty_to_invoice', 'price_unit', 'price_subtotal', 'invoice_status'], limit=100000, context=CTX)
byo = collections.defaultdict(list)
for l in lines:
    byo[l['order_id'][0]].append(l)
seul = [o for o in orders if byo[o['id']] and all(l['product_id'] and l['product_id'][0] in pids for l in byo[o['id']])]
mixte = [o for o in orders if byo[o['id']] and any(l['product_id'] and l['product_id'][0] in pids for l in byo[o['id']]) and o not in seul]
print('commandes confirmées jamais facturées (SARL MAQUIGNON) : %d | uniquement transport : %d (%.2f € HT) | mixtes avec transport : %d (%.2f € HT)' % (
    len(orders), len(seul), sum(o['amount_untaxed'] for o in seul), len(mixte), sum(o['amount_untaxed'] for o in mixte)))
today = datetime.date.today()
tache_ids = sorted({o['x_studio_tache'][0] for o in seul if o['x_studio_tache']})
taches = {t['id']: t for t in x('project.task', 'read', tache_ids, ['name', 'project_id', 'stage_id', 'state', 'x_studio_chauffeur', 'planned_date_begin', 'x_studio_ca_transport'] if False else ['name', 'project_id', 'stage_id', 'state', 'planned_date_begin'], context=CTX)} if tache_ids else {}
pick_ids = sorted({pk for o in seul for pk in o['picking_ids']})
picks = {pk['id']: pk['state'] for pk in x('stock.picking', 'read', pick_ids, ['state'], context=CTX)} if pick_ids else {}


def resume(o):
    fams = collections.Counter(famille(l['product_id'][1]) for l in byo[o['id']])
    return ', '.join('%s ×%d' % (k, v) for k, v in fams.most_common())


rows = []
for o in sorted(seul, key=lambda q: q['date_order']):
    t = taches.get(o['x_studio_tache'][0]) if o['x_studio_tache'] else None
    ls = byo[o['id']]
    rows.append({
        'Commande': o['name'], 'Date': o['date_order'][:10], 'Client': o['partner_id'][1], 'Montant HT': round(o['amount_untaxed'], 2),
        'Âge (j)': (today - datetime.date.fromisoformat(o['date_order'][:10])).days, 'Statut facturation': o['invoice_status'],
        'À facturer (case)': 'oui' if o['x_studio_a_facturer'] else 'non', 'Familles': resume(o), 'Nb lignes': len(ls),
        'Articles': ' | '.join('%s ×%g à %.2f' % (l['product_id'][1][:40], l['product_uom_qty'], l['price_unit']) for l in ls[:4]),
        'Origine': o['origin'] or '', 'Réf client': o['client_order_ref'] or '', 'Créée par': (o['create_uid'] or ['', ''])[1], 'Créée le': o['create_date'][:10], 'Vendeur': (o['user_id'] or ['', ''])[1],
        'Tâche transport': t and t['name'][:50] or '', 'Projet': t and t['project_id'] and t['project_id'][1] or '', 'Étape tâche': t and t['stage_id'] and t['stage_id'][1] or '',
        'Tâche prévue le': t and t['planned_date_begin'] and t['planned_date_begin'][:10] or '', 'Adresse transport': o['x_studio_transport_adresse'] or '',
        'BL': ', '.join(sorted(set(picks[pk] for pk in o['picking_ids'] if pk in picks))), 'Commentaire': (o['x_studio_transport_commentaire'] or '')[:80],
    })
print('\npar mois :', dict(sorted(collections.Counter(r['Date'][:7] for r in rows).items())))
print('par famille dominante :', dict(collections.Counter(r['Familles'].split(' ×')[0] for r in rows).most_common()))
print('par créateur :', dict(collections.Counter(r['Créée par'] for r in rows).most_common(6)))
print('avec tâche de transport liée : %d | projets : %s' % (len([r for r in rows if r['Tâche transport']]), dict(collections.Counter(r['Projet'] for r in rows if r['Projet']).most_common(4))))
print('case « À facturer » cochée : %d' % len([r for r in rows if r['À facturer (case)'] == 'oui']))
print('montant par client (top 12) :')
mt = collections.defaultdict(float); nb = collections.Counter()
for r in rows:
    mt[r['Client']] += r['Montant HT']; nb[r['Client']] += 1
for k, v in sorted(mt.items(), key=lambda kv: -kv[1])[:12]:
    print('   %-40s %3d BC  %10.2f € HT' % (k[:40], nb[k], v))
print('exemples :')
for r in rows[:5] + rows[-3:]:
    print('   %s %s %-28s %9.2f  %-38s tâche %s' % (r['Commande'], r['Date'], r['Client'][:28], r['Montant HT'], r['Familles'][:38], r['Tâche transport'][:30]))
from openpyxl import Workbook
wb = Workbook(); ws = wb.active; ws.title = 'Commandes'
cols = list(rows[0].keys()) if rows else []
ws.append(cols)
for r in rows:
    ws.append([r[k] for k in cols])
ws2 = wb.create_sheet('Synthèse')
ws2.append(['Mois', 'Nb BC', 'Montant HT'])
for mo_ in sorted({r['Date'][:7] for r in rows}):
    ws2.append([mo_, len([r for r in rows if r['Date'][:7] == mo_]), round(sum(r['Montant HT'] for r in rows if r['Date'][:7] == mo_), 2)])
ws2.append([]); ws2.append(['Client', 'Nb BC', 'Montant HT'])
for k, v in sorted(mt.items(), key=lambda kv: -kv[1]):
    ws2.append([k, nb[k], round(v, 2)])
ws3 = wb.create_sheet('Mixtes (non listées)'); ws3.append(['Commande', 'Date', 'Client', 'Montant HT', 'Lignes transport', 'Lignes autres'])
for o in sorted(mixte, key=lambda q: q['date_order']):
    ls = byo[o['id']]
    ws3.append([o['name'], o['date_order'][:10], o['partner_id'][1], round(o['amount_untaxed'], 2), len([l for l in ls if l['product_id'][0] in pids]), len([l for l in ls if l['product_id'][0] not in pids])])
chemin = 'C:/Users/xavfe/Desktop/Maquignon/BC_transport_jamais_factures_%s%s.xlsx' % (today.isoformat(), ('_avant_' + AVANT) if AVANT else '')
wb.save(chemin); print('\nExcel :', chemin)

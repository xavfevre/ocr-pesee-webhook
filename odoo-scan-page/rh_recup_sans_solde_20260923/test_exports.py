# -*- coding: utf-8 -*-
"""Exports sur la base de test : feuille hebdo JOLLY septembre (vs feuille papier), export paie, Silae."""
import os, ssl, sys, importlib.util, xmlrpc.client, io
sys.stdout.reconfigure(encoding='utf-8')
U, D = 'https://testmaq230926v2.odoo.com', 'testmaq230926v2'
us = os.environ['ODOO_USER']; p = os.environ['ODOO_PWD']
c = ssl.create_default_context()
uid = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/common', context=c).authenticate(D, us, p, {})
m = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/object', context=c)
def call(mo, me, *a, **k): return m.execute_kw(D, uid, p, mo, me, list(a), k)
sys.path.insert(0, 'ocr')
spec = importlib.util.spec_from_file_location('ex', 'ocr/export_heures.py'); ex = importlib.util.module_from_spec(spec); spec.loader.exec_module(ex)
import openpyxl
emp = call('hr.employee', 'search_read', [['name', '=', 'JOLLY Floran']], ['id'])[0]['id']
data = ex.build_feuille(call, '2026-09', emp)
open('captures_rh/feuille_jolly_2026-09.xlsx', 'wb').write(data)
wb = openpyxl.load_workbook(io.BytesIO(data))
ws = wb['Récap']
print('feuille JOLLY : onglets', wb.sheetnames)
for r in range(11, 60):
    vals = [ws.cell(r, cc).value for cc in range(2, 12)]
    if any(v is not None for v in vals):
        print('  ', r, [(v.strftime('%a %d/%m') if hasattr(v, 'strftime') and not hasattr(v, 'hour') else (v.strftime('%H:%M') if hasattr(v, 'hour') and not hasattr(v, 'year') else v)) for v in vals])
data2 = ex.build(call, '2026-09', 1)
open('captures_rh/export_paie_2026-09.xlsx', 'wb').write(data2)
wb2 = openpyxl.load_workbook(io.BytesIO(data2))
ws2 = wb2['JOLLY Floran']
print('export paie JOLLY : en-têtes', [ws2.cell(4, j).value for j in range(1, 20)])
print('                    valeurs ', [ws2.cell(5, j).value for j in range(1, 20)])
for r in range(7, 60):
    if ws2.cell(r, 2).value and str(ws2.cell(r, 2).value)[:2] in ('11', '14', '18'):
        print('  ', [ws2.cell(r, j).value for j in range(1, 12)])
data3 = ex.build_silae(call, '2026-09', 1)
wb3 = openpyxl.load_workbook(io.BytesIO(data3))
print('silae EVP lignes :', [[wb3['EVP'].cell(r, j).value for j in range(1, 6)] for r in range(2, wb3['EVP'].max_row + 1) if 'JOLLY' in str(wb3['EVP'].cell(r, 2).value)])
print('silae absences :', [[wb3['Absences'].cell(r, j).value for j in range(1, 9)] for r in range(2, wb3['Absences'].max_row + 1) if 'JOLLY' in str(wb3['Absences'].cell(r, 2).value)])

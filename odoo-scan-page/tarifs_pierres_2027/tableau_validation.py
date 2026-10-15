# -*- coding: utf-8 -*-
"""Tableau de validation des tarifs pierres 2027 (15/10/2026, demande Xavier : « un tableau simple des prix actuels avec les
phases d'augmentation que je valide », puis « ajoute une colonne pour ces clients [à prix négociés] »). Une ligne par pierre,
dans l'unité du devis (m² pour les pierres d'épaisseur, m³ sinon) :
  PRO liste Odoo 2026 | PRO facturé 2026 | PRO 01/11/2026 | PRO 01/01/2027 | hausse | PARTICULIER actuel (fiche) | 01/11 | 01/01 | hausse |
  une colonne par client à prix négociés (prix de sa liste Odoo en gras, sinon moyenne facturée 2026 en gris) | Validé | Remarque.
  Les colonnes jaunes sont celles qui seront chargées ; le chargeur lit ce fichier (feuille « Validation », colonnes repérées par leur titre).
  -> Desktop/Maquignon/Tarifs_pierres_2027_validation_2026-10-15.xlsx   (lecture seule côté Odoo)"""
import os, sys, re, xmlrpc.client
from collections import defaultdict
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
ICI = os.path.dirname(os.path.abspath(__file__))
SORTIE = r"C:\Users\xavfe\Desktop\Maquignon\Tarifs_pierres_2027_validation_2026-10-15.xlsx"
CLIENTS = [(15893, 'Lefèvre 44'), (15891, 'Lefèvre 56'), (15890, 'Lefèvre Alençon'), (15758, 'Crocodile'), (15746, 'Cobat 49'), (35816, 'BPNR 37'),
           (15671, "Pierre d'Anjou"), (15721, 'Cazy'), (15803, 'Fonteneau'), (15933, 'Th. Moreau'), (21539, 'Vermorel')]
U = 'https://maquignon.odoo.com'; D = 'maquignon'; us = os.environ['ODOO_USER']; p = os.environ['ODOO_PWD']
uid = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/common').authenticate(D, us, p, {})
m = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/object')
CTX = {'allowed_company_ids': [1], 'company_id': 1, 'active_test': False}
x = lambda mo, me, *a, **k: m.execute_kw(D, uid, p, mo, me, list(a), dict(k, context=dict(CTX, **k.pop('context', {}))))
nombre = lambda v: v if isinstance(v, (int, float)) else None
w = load_workbook(os.path.join(ICI, 'v3.xlsx'), data_only=False)['Grille 2027 (2)']
prods = x('product.product', 'search_read', [['product_tmpl_id.company_id', 'in', [1, False]]], fields=['default_code', 'display_name', 'lst_price', 'active'], limit=5000)
by_code = {pr['default_code']: pr for pr in prods if pr['default_code']}
nom_net = lambda s: re.sub(r'^\[[^\]]+\] ', '', s).replace(' (archivé)', '').strip()
by_name = {}
for pr in prods:
    by_name.setdefault(nom_net(pr['display_name']), []).append(pr)


def apparier(art):
    mm = re.match(r'\[([^\]]+)\] ', art)
    if mm and mm.group(1) in by_code:
        return by_code[mm.group(1)]
    c = by_name.get(nom_net(art), []); c = [pr for pr in c if pr['active']] or c
    return c[0] if len(c) == 1 else None


# Tarif Pro 2026 (liste 299) : règles par variante (pas de règle générale -> repli sur le prix fiche)
regles = x('product.pricelist.item', 'search_read', [['pricelist_id', '=', 299], ['applied_on', '=', '0_product_variant']], fields=['product_id', 'compute_price', 'fixed_price'])
r_var = {r['product_id'][0]: r for r in regles if r['product_id'] and r['compute_price'] == 'fixed'}

rows, cat = [], None
for r in range(2, w.max_row + 1):
    a = w.cell(r, 1).value
    if isinstance(a, str) and '›' in a:
        cat = a; continue
    if not a or not cat or not cat.startswith('PIERRES'):
        continue
    unite = (w.cell(r, 2).value or '').strip(); K, L, M, O = (w.cell(r, c).value for c in (11, 12, 13, 15))
    if isinstance(L, str) and L.startswith('='):
        L = (M / 1.1) if ('/1.1' in L and nombre(M)) else (K * 1.03 if nombre(K) else None)
    pr = apparier(a)
    if not pr or not nombre(O) or O <= 1 or not nombre(M) or M <= 1 or unite == 'Unité':
        continue
    ep = re.search(r'\((\d+)\s*cm\)', a); ep = int(ep.group(1)) / 100.0 if (unite == 'm²' and ep) else None
    if unite == 'm²' and not ep:
        continue
    coef = ep or 1.0
    pro_m3 = r_var[pr['id']]['fixed_price'] if pr['id'] in r_var else pr['lst_price']
    code = re.match(r'\[([^\]]+)\]', a)
    rows.append({'pid': pr['id'], 'coef': coef, 'cat': cat.split('›')[-1].strip(), 'code': code.group(1) if code else '', 'art': nom_net(a), 'unite': unite,
                 'pro_liste': round(pro_m3 * coef, 2), 'pro1': round(L), 'pro2': round(M),
                 'part_act': round(pr['lst_price'] * coef, 2), 'part1': round(O / 1.1), 'part2': round(O)})
pids = [q['pid'] for q in rows]; coef_de = {q['pid']: q['coef'] for q in rows}

# prix réellement facturés en 2026 (société 1), dans l'unité du devis : par liste (pro) et par client négocié
lines = x('account.move.line', 'search_read', [['move_id.move_type', '=', 'out_invoice'], ['move_id.state', '=', 'posted'], ['move_id.invoice_date', '>=', '2026-01-01'], ['product_id', 'in', pids], ['company_id', '=', 1]], fields=['product_id', 'price_unit', 'quantity', 'discount', 'partner_id'], limit=20000)
part_ids = list({l['partner_id'][0] for l in lines})
plist = {pp['id']: (pp['commercial_partner_id'] and pp['commercial_partner_id'][0]) for pp in x('res.partner', 'read', part_ids, ['commercial_partner_id'])}
liste_de = {pp['id']: (pp['specific_property_product_pricelist'] and pp['specific_property_product_pricelist'][0]) for pp in x('res.partner', 'read', list(set(plist.values())), ['specific_property_product_pricelist'])}
acc = defaultdict(lambda: [0.0, 0.0, 0]); acc_cli = defaultdict(lambda: [0.0, 0.0, 0])
for l in lines:
    if l['quantity'] <= 0:
        continue
    cp = plist.get(l['partner_id'][0]); net = l['price_unit'] * (1 - (l['discount'] or 0) / 100.0)
    if liste_de.get(cp) == 299:
        a_ = acc[l['product_id'][0]]; a_[0] += net * l['quantity']; a_[1] += l['quantity']; a_[2] += 1
    if cp in dict(CLIENTS):
        a_ = acc_cli[(cp, l['product_id'][0])]; a_[0] += net * l['quantity']; a_[1] += l['quantity']; a_[2] += 1
for q in rows:
    a_ = acc.get(q['pid']); q['pro_fact'] = round(a_[0] / a_[1], 2) if a_ and a_[1] else None; q['pro_fact_n'] = a_[2] if a_ else 0
# listes propres des clients négociés (prix de liste au m³ -> unité du devis)
cli_liste = {}
for cid, nom in CLIENTS:
    pp = x('res.partner', 'read', [cid], ['name', 'specific_property_product_pricelist'])[0]
    lid = pp['specific_property_product_pricelist'] and pp['specific_property_product_pricelist'][0]
    if lid and lid not in (299, 302, 12):
        for r in x('product.pricelist.item', 'search_read', [['pricelist_id', '=', lid], ['applied_on', '=', '0_product_variant'], ['compute_price', '=', 'fixed'], ['product_id', 'in', pids]], fields=['product_id', 'fixed_price', 'date_start', 'date_end']):
            if not r['date_start']:
                cli_liste[(cid, r['product_id'][0])] = round(r['fixed_price'] * coef_de[r['product_id'][0]], 2)
        print('%-16s liste propre %s : %d prix pierres' % (nom, pp['specific_property_product_pricelist'][1], sum(1 for k in cli_liste if k[0] == cid)))
    else:
        print('%-16s sur %s' % (nom, (pp['specific_property_product_pricelist'] and pp['specific_property_product_pricelist'][1]) or 'aucune liste'))

wb = Workbook(); ws = wb.active; ws.title = 'Validation'
gras = Font(bold=True); titre = Font(bold=True, size=13); gris = Font(color='666666', italic=True); grasbleu = Font(bold=True, color='1D4ED8')
bleu = PatternFill('solid', fgColor='EEF2FF'); jaune = PatternFill('solid', fgColor='FFF8DC'); vertc = PatternFill('solid', fgColor='E8F5E9'); mauve = PatternFill('solid', fgColor='F3E8FF')
fin = Side(style='thin', color='CBD5E1'); bord = Border(bottom=fin)
NCOL = 12 + len(CLIENTS) + 2
ws.append(['Tarifs pierres 2027 : tableau de validation (SARL Maquignon)']); ws['A1'].font = titre
ws.append(["Prix HT dans l'unité du devis : m² pour les pierres d'épaisseur, m³ pour massifs, tranches et blocs, tonne pour l'enrochement. "
           "PRO liste = Tarif Pro 2026 d'Odoo (prix fiche s'il n'y a pas de règle) ; PRO facturé = moyenne des factures 2026 aux clients pro (nombre de lignes entre parenthèses) ; PARTICULIER actuel = prix fiche. "
           "Colonnes jaunes = ce qui sera chargé dans Odoo (corrigez directement si besoin) ; colonne verte = votre validation. "
           "Colonnes mauves = clients à prix négociés, non chargés : en bleu gras le prix de leur liste Odoo, en gris la moyenne de ce qu'ils ont payé en 2026 (nombre de lignes entre parenthèses)."])
ws['A2'].font = gris; ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=NCOL); ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 52
ws.append([])
tete = ['Code', 'Pierre', 'Unité', 'PRO liste Odoo 2026', 'PRO facturé 2026', 'PRO 1er nov. 2026', 'PRO 1er janv. 2027', 'Hausse pro 1er janv. (vs facturé, sinon liste)',
        'PARTICULIER actuel (fiche)', 'PARTICULIER 1er nov. 2026', 'PARTICULIER 1er janv. 2027', 'Hausse particulier 1er janv.'] + [n for _, n in CLIENTS] + ['Validé (oui / non)', 'Remarque']
ws.append(tete)
for c in ws[4]:
    c.font = gras; c.fill = bleu; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center'); c.border = bord
for i in range(len(CLIENTS)):
    ws.cell(4, 13 + i).fill = mauve
ws.row_dimensions[4].height = 44
cat_cour = None
for q in rows:
    if q['cat'] != cat_cour:
        cat_cour = q['cat']; ws.append([cat_cour]); ws.cell(ws.max_row, 1).font = gras; ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=NCOL)
    ln = ws.max_row + 1
    fact = ('%.2f (%d)' % (q['pro_fact'], q['pro_fact_n'])) if q['pro_fact'] else ''
    base_pro = q['pro_fact'] or q['pro_liste']
    cellules = [q['code'], q['art'], q['unite'], q['pro_liste'], fact, q['pro1'], q['pro2'], (q['pro2'] / base_pro - 1) if base_pro else None,
                q['part_act'], q['part1'], q['part2'], (q['part2'] / q['part_act'] - 1) if q['part_act'] else None]
    for cid, _ in CLIENTS:
        if (cid, q['pid']) in cli_liste:
            cellules.append(cli_liste[(cid, q['pid'])])
        elif (cid, q['pid']) in acc_cli:
            a_ = acc_cli[(cid, q['pid'])]; cellules.append('%.2f (%d)' % (a_[0] / a_[1], a_[2]))
        else:
            cellules.append('')
    cellules += ['', '']
    ws.append(cellules)
    for col in (6, 7, 10, 11):
        ws.cell(ln, col).fill = jaune
    for col in (8, 12):
        ws.cell(ln, col).number_format = '+0 %;-0 %;0 %'; ws.cell(ln, col).font = gris
    for col in (4, 6, 7, 9, 10, 11):
        ws.cell(ln, col).number_format = '# ##0.00'
    ws.cell(ln, 5).font = gris; ws.cell(ln, 5).alignment = Alignment(horizontal='right')
    for i, (cid, _) in enumerate(CLIENTS):
        c = ws.cell(ln, 13 + i); c.fill = mauve; c.alignment = Alignment(horizontal='right')
        if (cid, q['pid']) in cli_liste:
            c.font = grasbleu; c.number_format = '# ##0.00'
        else:
            c.font = gris
    ws.cell(ln, 13 + len(CLIENTS)).fill = vertc
    for col in range(1, NCOL + 1):
        ws.cell(ln, col).border = bord
largeurs = [13, 42, 6, 10, 12, 10, 10, 11, 11, 11, 11, 10] + [11] * len(CLIENTS) + [9, 26]
for i, wd in enumerate(largeurs):
    ws.column_dimensions[ws.cell(4, i + 1).column_letter].width = wd
ws.freeze_panes = 'C5'
ws.page_setup.orientation = 'landscape'; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0; ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.print_title_rows = '4:4'
wb.save(SORTIE)
print('Excel :', SORTIE, '|', len(rows), 'pierres |', len(CLIENTS), 'colonnes clients')
nb = sum(1 for q in rows for cid, _ in CLIENTS if (cid, q['pid']) in cli_liste or (cid, q['pid']) in acc_cli)
print('cases clients renseignées :', nb)

# -*- coding: utf-8 -*-
"""Point « prix fiche (= tarif public actuel) vs nouveaux tarifs particuliers » (15/10/2026, demande Xavier).
Pour chaque pierre chargée : unité de vente de la grille, prix particulier ACTUEL (prix fiche Odoo, ramené au m² pour
les pierres d'épaisseur = fiche × épaisseur, ce que fait le devis), vague 1 (01/11 = PARTICULIER 2027 ÷ 1,1), vague 2
(01/01 = PARTICULIER 2027), écarts ; pro 2026 / 01/11 / 2027 pour mémoire.
  -> Desktop/Maquignon/Point_tarifs_public_vs_fiche_2026-10-15.xlsx + résumé à l'écran. Lecture seule."""
import os, sys, re, xmlrpc.client
sys.stdout.reconfigure(encoding='utf-8')
from openpyxl import load_workbook, Workbook
from openpyxl.styles import Font, PatternFill, Alignment
ICI = os.path.dirname(os.path.abspath(__file__))
SORTIE = r"C:\Users\xavfe\Desktop\Maquignon\Point_tarifs_public_vs_fiche_2026-10-15.xlsx"
U = 'https://maquignon.odoo.com'; D = 'maquignon'; us = os.environ['ODOO_USER']; p = os.environ['ODOO_PWD']
uid = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/common').authenticate(D, us, p, {})
m = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/object')
CTX = {'allowed_company_ids': [1], 'company_id': 1, 'active_test': False}
x = lambda mo, me, *a, **k: m.execute_kw(D, uid, p, mo, me, list(a), dict(k, context=dict(CTX, **k.pop('context', {}))))
nombre = lambda v: v if isinstance(v, (int, float)) else None
w = load_workbook(os.path.join(ICI, 'v3.xlsx'), data_only=False)['Grille 2027 (2)']
prods = x('product.product', 'search_read', [['product_tmpl_id.company_id', 'in', [1, False]]], fields=['default_code', 'display_name', 'lst_price', 'active'], limit=5000)
by_code = {pr['default_code']: pr for pr in prods if pr['default_code']}
nom_net = lambda s: re.sub(r'^\[[^\]]+\] ', '', s).replace(' (archivé)', '').strip()
by_name = {}
for pr in prods:
    by_name.setdefault(nom_net(pr['display_name']), []).append(pr)


def apparier(art):
    mm = re.match(r'\[([^\]]+)\] ', art)
    if mm and mm.group(1) in by_code:
        return by_code[mm.group(1)]
    c = by_name.get(nom_net(art), []); c = [pr for pr in c if pr['active']] or c
    return c[0] if len(c) == 1 else None


rows, cat = [], None
for r in range(2, w.max_row + 1):
    a = w.cell(r, 1).value
    if isinstance(a, str) and '›' in a:
        cat = a; continue
    if not a or not cat or not cat.startswith('PIERRES'):
        continue
    unite = (w.cell(r, 2).value or '').strip(); K, L, M, O = (w.cell(r, c).value for c in (11, 12, 13, 15))
    if isinstance(L, str) and L.startswith('='):
        L = (M / 1.1) if ('/1.1' in L and nombre(M)) else (K * 1.03 if nombre(K) else None)
    pr = apparier(a)
    if not pr or not nombre(O) or O <= 1 or unite == 'Unité':
        continue
    ep = re.search(r'\((\d+)\s*cm\)', a); ep = int(ep.group(1)) / 100.0 if (unite == 'm²' and ep) else None
    if unite == 'm²' and not ep:
        continue
    actuel = pr['lst_price'] * (ep or 1.0)        # ce que paie un particulier aujourd'hui (le devis fait fiche × épaisseur)
    v1, v2 = round(O / 1.1), round(O)
    rows.append({'cat': cat.split('›')[-1].strip(), 'art': a, 'unite': unite, 'actuel': round(actuel, 2), 'v1': v1, 'v2': v2,
                 'd1': (v1 / actuel - 1) * 100 if actuel else None, 'd2': (v2 / actuel - 1) * 100 if actuel else None,
                 'k': nombre(K), 'l': round(L) if nombre(L) else None, 'm': round(M) if nombre(M) else None,
                 'dpro': ((M / K) - 1) * 100 if (nombre(K) and nombre(M)) else None})
wb = Workbook(); ws = wb.active; ws.title = 'Public vs fiche'
tete = ['Catégorie', 'Article', 'Unité', 'Prix particulier actuel (fiche)', 'Particulier 01/11/2026', 'Écart 01/11 %', 'Particulier 01/01/2027', 'Écart 01/01 % (cumulé)', 'Pro 2026', 'Pro 01/11/2026', 'Pro 2027', 'Hausse pro 2027 %']
ws.append(tete)
for c in ws[1]:
    c.font = Font(bold=True); c.fill = PatternFill('solid', fgColor='EEF2FF'); c.alignment = Alignment(wrap_text=True, vertical='top')
rouge = PatternFill('solid', fgColor='FDE2E2'); orange = PatternFill('solid', fgColor='FFF4E0'); vert = PatternFill('solid', fgColor='E3F6E5')
for q in rows:
    ws.append([q['cat'], q['art'], q['unite'], q['actuel'], q['v1'], round(q['d1'], 1) if q['d1'] is not None else None, q['v2'], round(q['d2'], 1) if q['d2'] is not None else None, q['k'], q['l'], q['m'], round(q['dpro'], 1) if q['dpro'] is not None else None])
    for col, d in ((6, q['d1']), (8, q['d2'])):
        cell = ws.cell(ws.max_row, col)
        if d is None:
            continue
        cell.fill = rouge if d >= 20 else (orange if d >= 10 else (vert if d < 0 else PatternFill()))
for col, wd in zip('ABCDEFGHIJKL', (14, 52, 7, 16, 14, 11, 14, 14, 10, 12, 10, 12)):
    ws.column_dimensions[col].width = wd
ws.freeze_panes = 'C2'
wb.save(SORTIE)
print('Excel :', SORTIE, '|', len(rows), 'pierres')
# résumé
d1 = [q for q in rows if q['d1'] is not None]
print('\nVague 1 (01/11) pour les particuliers, par rapport au prix fiche actuel : médiane %+.1f %%, de %+.1f %% à %+.1f %%' % (sorted(q['d1'] for q in d1)[len(d1) // 2], min(q['d1'] for q in d1), max(q['d1'] for q in d1)))
print('Vague 2 (01/01) cumulée : médiane %+.1f %%, de %+.1f %% à %+.1f %%' % (sorted(q['d2'] for q in d1)[len(d1) // 2], min(q['d2'] for q in d1), max(q['d2'] for q in d1)))
print('  baisse ou stable au 01/11 : %d pierres | +0 à +10 %% : %d | +10 à +20 %% : %d | plus de +20 %% : %d' % (sum(1 for q in d1 if q['d1'] <= 0.5), sum(1 for q in d1 if 0.5 < q['d1'] < 10), sum(1 for q in d1 if 10 <= q['d1'] < 20), sum(1 for q in d1 if q['d1'] >= 20)))
print('\nLes plus fortes hausses particulier au 01/11 (prix actuel -> 01/11 -> 01/01) :')
for q in sorted(d1, key=lambda q: -q['d1'])[:12]:
    print('   %-60s %-3s %9.2f -> %6d (%+5.1f %%) -> %6d (%+5.1f %%)' % (q['art'][:60], q['unite'][:3], q['actuel'], q['v1'], q['d1'], q['v2'], q['d2']))
print('\nBaisses ou quasi-stables au 01/11 :')
for q in sorted(d1, key=lambda q: q['d1'])[:8]:
    print('   %-60s %-3s %9.2f -> %6d (%+5.1f %%) -> %6d (%+5.1f %%)' % (q['art'][:60], q['unite'][:3], q['actuel'], q['v1'], q['d1'], q['v2'], q['d2']))

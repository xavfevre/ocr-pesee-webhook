# -*- coding: utf-8 -*-
"""Ajoute un onglet « Remarques » (revue de cohérence du 17/09/2026) au fichier Tarifs 2027, sans toucher aux prix."""
import sys, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
sys.stdout.reconfigure(encoding='utf-8')
SRC = 'C:/Users/xavfe/Desktop/Maquignon/Tarifs_2027_Maquignon_fiche_mini15_tuffeau1500_ecartfixe_entier_categories_2026-09-17.xlsx'
OUT = SRC.replace('_categories_', '_categories_remarques_')
wb = openpyxl.load_workbook(SRC)
if 'Remarques' in wb.sheetnames:
    del wb['Remarques']
ws = wb.create_sheet('Remarques', 1)
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='01666B'); T = PatternFill('solid', fgColor='CCE5E4'); ROUGE = PatternFill('solid', fgColor='F8CBAD'); JAUNE = PatternFill('solid', fgColor='FFF2CC')
cols = ['Sujet', 'Article(s)', 'Aujourd\'hui dans la grille (pro / particulier)', 'Ce qui cloche', 'Proposition', 'Priorité']
ws.append(cols)
for i, cell in enumerate(ws[1], 1):
    cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical='center')
for col, w in zip('ABCDEF', (22, 46, 30, 52, 44, 9)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'


def titre(t):
    ws.append([t]); ws.cell(ws.max_row, 1).font = Font(bold=True)
    for j in range(1, 7):
        ws.cell(ws.max_row, j).fill = T


def ligne(sujet, art, auj, pb, prop, prio):
    ws.append([sujet, art, auj, pb, prop, prio])
    r = ws.max_row
    for j in range(1, 7):
        ws.cell(r, j).alignment = Alignment(wrap_text=True, vertical='top')
    ws.cell(r, 6).fill = ROUGE if prio == 'Haute' else JAUNE


titre('CE QUI TIENT : blocs < tranches < pré-sciées ; prix au m³ qui monte quand l\'épaisseur baisse ; hiérarchie des pierres (tuffeau 1 275 < Migné 1 437 < Sireuil 1 565 < Tervoux 1 768 < Haims 1 891 < Thénac 2 145 < Richemont 2 664 €/m³ pro) ; transports à la tonne cohérents entre eux (4,10 à 7,15 €/t selon le départ) ; granulats et terre sans surprise.')
titre('PIERRES : incohérences entre variantes (comparaison au m³ = prix au m² ÷ épaisseur)')
ligne('Épaisseurs inversées', 'Tuffeau pré-scié 4 cm [TUF0004-PS]', '83 / 98 €/m²', 'Plus cher au m² que le 5 cm (82 / 96). Une dalle plus fine ne peut pas coûter plus qu\'une plus épaisse.', '79 / 93 €/m² (entre le 3 cm à 74 et le 5 cm à 82)', 'Haute')
ligne('Épaisseurs inversées', 'Haims pré-scié 4 cm [HAIMS0004-PS]', '138 / 162 €/m²', 'Au-dessus du 5 cm (129 / 152) et loin du 3 cm (108 / 127).', '120 / 141 €/m²', 'Haute')
ligne('Épaisseurs inversées', 'Migné pré-scié 3 cm [MIGNE0003-PS]', '118 / 139 €/m²', 'Plus cher que le 5 cm (115 / 135).', '108 / 127 €/m²', 'Haute')
ligne('Épaisseurs égales', 'Tervoux pré-scié 2 cm et 3 cm', '108 / 127 €/m² pour les deux', 'Même prix pour deux épaisseurs (le 2 cm devrait être un peu moins cher au m²).', '2 cm à 98 / 115 €/m²', 'Moyenne')
ligne('Tranche plus chère que le pré-scié', 'Haims tranche 10 cm [HAIMS0010-TR]', '286 / 336 €/m² (2 860 €/m³)', 'Une tranche 4 faces au-dessus du pré-scié 6 faces de même épaisseur (213 / 251). Les tranches 11-20 cm sont à 1 853 €/m³.', '190 / 224 €/m²', 'Haute')
ligne('Bande 4 faces plus chère que le pré-scié', 'Haims bande 4 faces 10 cm [HAIMS0010-BA]', '215 / 253 €/m²', 'Au-dessus du pré-scié 6 faces 10 cm (213 / 251) ; les autres bandes Haims sont bien en dessous des pré-sciés.', '200 / 235 €/m²', 'Moyenne')
ligne('Doublons « (U) »', 'Tuffeau pré-scié 6 faces (U) massif ; bandes 4 faces (U) massif, 3, 5, 7, 8, 9 cm', 'massif (U) 1 363 / 1 604 contre standard 1 275 / 1 500 ; bande 8 cm (U) 100 / 118 contre 111 / 131 ; bande massif (U) 1 004 / 1 181 contre 1 081 / 1 272', 'Deux articles pour le même produit avec des prix différents (la variante (U) n\'a pas reçu le 1 500 imposé).', 'Aligner les (U) sur le standard, ou me dire ce que signifie (U) si c\'est une autre qualité', 'Haute')
ligne('Variantes client', 'Tuffeau pré-scié 6 faces (COULMEAU), 3 variantes au m²', '744 / 875, 296 / 348, 245 / 288 €/m² soit 980 à 1 078 €/m³', 'Sous le prix du massif (1 275 €/m³). Ressemble à des prix négociés pour un client mis en article.', 'À vérifier : prix client (onglet client) ou vraie variante ?', 'Moyenne')
ligne('Ordre des épaisseurs', 'Richemont pré-scié 7 et 8 cm', '7 cm 228 / 268, 8 cm 277 / 326 €/m²', 'Au m³, le 7 cm (3 257) ressort sous le 8 cm (3 462) : pas grave au m², mais peu de factures (1 client).', 'Laisser, à revoir quand il y aura des ventes', 'Basse')
titre('ARTICLES QUI N\'ONT PAS LEUR PLACE DANS UNE GRILLE DE PRIX FIXES (prix au devis)')
ligne('Prestations au devis', 'Taille de pierre (Forfait) ; TP (Forfait) ; TP (Prix Forfaitaire) ; TP autoliquidée ; Location de matériel TP (Forfait)', '322 / 379 ; 880 / 1 035 ; 440 / 518 ; 820 / 965 ; 880', 'Prix facturés de 16 à 3 310 € pour le même article : la moyenne ne veut rien dire, c\'est un devis à chaque fois.', 'Ne pas mettre de prix dans les listes (prix saisi sur le devis)', 'Haute')
ligne('Prix à 1 €', 'Richemont à l\'unité ; Balustres en Tervoux ; Vente produit TVA 0 % ; Transport divers (Tonne) ; Transport retour de marchandises (Tonne)', '1 / 1', 'Prix fiche à 1 € = prix saisi à la commande.', 'Hors listes, ou vrai prix si tu en as un', 'Moyenne')
ligne('Surcharge carburant', 'Indexation Gasoil', '79 / 93 € l\'unité', 'Une indexation carburant se gère en pourcentage du transport, pas en prix fixe.', 'Règle en % (par ex. +x % sur les lignes transport) ou hors listes', 'Moyenne')
ligne('Articles archivés', 'Transport de granulats (Tonne) (archivé) ; Transport de granulats (Forfait) (archivé) ; Location semi fond-mouvant sans catégorie', '6,60 ; 150 ; 407', 'Encore facturés en 2026 alors qu\'archivés.', 'Ne pas charger ; vérifier que les commandes utilisent les articles actifs', 'Basse')
titre('UNITÉS FAUSSES SUR LES FICHES ARTICLES')
ligne('Forfait vendu à la tonne', 'Transport de granulats (Forfait) ; Transport de terre de gobetage (Forfait)', '92,55 €/t ; 447,55 €/t', 'Unité « Tonne » sur un forfait : prix avec centimes et quantité saisie en tonnes.', 'Passer l\'unité en Forfait (prix 93 et 448 €)', 'Haute')
ligne('Unité « Hours » sur des forfaits', 'Transfert de matériel (Forfait) ; Location de matériel TP (Forfait)', '781 ; 880 par « heure »', 'Le forfait est vendu à l\'heure dans Odoo.', 'Unité Forfait', 'Moyenne')
ligne('Unité « Unité » sur des transports', 'Transport de pierres (Forfait Blocs / Palettes / Tranches) ; Transport retour de marchandises (Forfait)', '583 ; 385 ; 491 ; 418', 'Cohérent entre eux, mais l\'unité devrait être Forfait pour être lisible sur les devis.', 'Unité Forfait', 'Basse')
titre('GRANULATS, TRANSPORT, LOCATIONS : rien de choquant')
ligne('Négoce graviers', 'Autres Graviers, Graviers Haims, Béton recyclé', '21,70 à 39,40 €/t pro ; béton recyclé 8,80', 'Dans la fourchette du marché livré (25-35 €/t concassé). Départ carrière les concurrents affichent 13-15 € HT/t.', 'RAS', 'Basse')
ligne('Locations à la journée', 'Benne (814) ; Semi-remorque (792) ; Camion 8x4 (726) ; Camion 6x4 (583)', 'voir grille', 'La benne à la journée au-dessus de la semi-remorque : à confirmer.', 'Vérifier le prix benne', 'Basse')
ligne('Location au tour', 'Duo au tour (484) ; Solo au tour (154)', 'voir grille', 'Duo à 3 fois le solo : peu de factures, à confirmer.', 'Vérifier', 'Basse')
ws.append([]); ws.append(['Revue faite le 17/09/2026 sur le fichier tarifs 2027 (écart fixe 15 %, tuffeau massif 1 500, arrondis à l\'euro). Aucun prix modifié dans les autres onglets.'])
wb.save(OUT)
print('fichier :', OUT, '| remarques :', ws.max_row - 1)

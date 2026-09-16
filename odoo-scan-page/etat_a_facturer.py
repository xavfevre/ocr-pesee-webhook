# -*- coding: utf-8 -*-
"""État des commandes clients « à facturer » (toutes sociétés) : une ligne par commande avec le reste à facturer,
l'ancienneté, factures et BL, et un classement ; un onglet avec le détail des lignes, un onglet de synthèse."""
import os, ssl, sys, xmlrpc.client, collections, datetime
sys.stdout.reconfigure(encoding='utf-8')
U = 'https://maquignon.odoo.com'; D = 'maquignon'; us = os.environ['ODOO_USER']; p = os.environ['ODOO_PWD']
c = ssl.create_default_context()
uid = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/common', context=c).authenticate(D, us, p, {})
m = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/object', context=c)
x = lambda mo, me, *a, **k: m.execute_kw(D, uid, p, mo, me, list(a), k)
ctx = {'allowed_company_ids': [1, 2, 3, 4, 13], 'active_test': False}
today = datetime.date.today()
orders = x('sale.order', 'search_read', [['state', '=', 'sale'], ['invoice_status', '=', 'to invoice']],
           fields=['name', 'partner_id', 'date_order', 'amount_untaxed', 'amount_total', 'invoice_ids', 'picking_ids',
                   'company_id', 'origin', 'client_order_ref', 'user_id', 'order_line'], limit=5000, context=ctx)
print('commandes « à facturer » :', len(orders), '| par société :', dict(collections.Counter(o['company_id'][1] for o in orders)))
oids = [o['id'] for o in orders]
lines = x('sale.order.line', 'search_read', [['order_id', 'in', oids], ['display_type', '=', False]],
          fields=['order_id', 'product_id', 'name', 'product_uom_qty', 'qty_delivered', 'qty_invoiced', 'qty_to_invoice',
                  'untaxed_amount_to_invoice', 'untaxed_amount_invoiced', 'price_subtotal', 'invoice_status', 'product_type'],
          limit=50000, context=ctx)
byo = collections.defaultdict(list)
for l in lines:
    byo[l['order_id'][0]].append(l)
pick_ids = sorted({pk for o in orders for pk in o['picking_ids']})
picks = {pk['id']: pk for pk in x('stock.picking', 'read', pick_ids, fields=['state', 'picking_type_code'], context=ctx)} if pick_ids else {}
inv_ids = sorted({i for o in orders for i in o['invoice_ids']})
invs = {i['id']: i for i in x('account.move', 'read', inv_ids, fields=['name', 'state', 'amount_untaxed', 'invoice_date', 'move_type'], context=ctx)} if inv_ids else {}


def famille(l):
    n = ((l['product_id'] and l['product_id'][1]) or l['name'] or '').lower()
    if 'transport' in n or 'location de mat' in n:
        return 'transport/location'
    if l['product_type'] == 'service':
        return 'service'
    return 'article'


rows = []
for o in orders:
    ls = [l for l in byo[o['id']] if l['invoice_status'] == 'to invoice']
    reste = sum(l['untaxed_amount_to_invoice'] for l in ls)
    fam = collections.Counter(famille(l) for l in ls)
    livre = sum(1 for l in ls if l['qty_delivered'] > 0)
    age = (today - datetime.date.fromisoformat(o['date_order'][:10])).days
    fact = [invs[i] for i in o['invoice_ids'] if i in invs and invs[i]['state'] == 'posted']
    bl_att = sum(1 for pk in o['picking_ids'] if picks.get(pk, {}).get('state') not in ('done', 'cancel'))
    bl_fait = sum(1 for pk in o['picking_ids'] if picks.get(pk, {}).get('state') == 'done')
    if reste < -0.005:
        cat = 'Sur-facturé (reste négatif) : avoir ou quantité à corriger'
    elif ls and set(fam) == {'transport/location'} and fact:
        cat = 'Transport/location seul, commande déjà facturée : à solder'
    elif ls and livre == len(ls):
        cat = 'Livré, non facturé : à facturer'
    elif age <= 30:
        cat = 'Récent (30 jours ou moins) : en cours'
    elif not fact and bl_fait == 0 and age > 60:
        cat = 'Ancienne, rien livré ni facturé : à vérifier (annuler ?)'
    elif fact:
        cat = 'Partiellement facturée, reste à facturer'
    else:
        cat = 'Non livrée, non facturée : à suivre'
    rows.append({'cde': o['name'], 'soc': o['company_id'][1], 'client': o['partner_id'][1], 'date': o['date_order'][:10], 'age': age,
                 'ref': o['client_order_ref'] or '', 'origine': o['origin'] or '', 'vendeur': (o['user_id'] and o['user_id'][1]) or '',
                 'total': round(o['amount_untaxed'], 2), 'reste': round(reste, 2), 'nb_lignes': len(ls),
                 'familles': ', '.join('%s (%d)' % (k, v) for k, v in fam.items()),
                 'factures': ', '.join(f['name'] for f in fact), 'montant_fact': round(sum(f['amount_untaxed'] for f in fact), 2),
                 'bl_attente': bl_att, 'bl_faits': bl_fait, 'cat': cat})

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
wb = Workbook(); ws = wb.active; ws.title = 'Commandes à facturer'
head = ['Catégorie', 'Société', 'Commande', 'Client', 'Date', 'Ancienneté (j)', 'Réf. client', 'Origine', 'Vendeur', 'Total cde HT',
        'Reste à facturer HT', 'Lignes à facturer', 'Familles', 'Factures existantes', 'Montant facturé HT', 'BL en attente', 'BL validés']
ws.append(head)
for r in sorted(rows, key=lambda r: (r['cat'], -r['reste'])):
    ws.append([r['cat'], r['soc'], r['cde'], r['client'], r['date'], r['age'], r['ref'], r['origine'], r['vendeur'], r['total'], r['reste'],
               r['nb_lignes'], r['familles'], r['factures'], r['montant_fact'], r['bl_attente'], r['bl_faits']])
for cell in ws[1]:
    cell.font = Font(bold=True, color='FFFFFF'); cell.fill = PatternFill('solid', fgColor='01666B')
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
for col, w in zip('ABCDEFGHIJKLMNOPQ', (46, 20, 12, 30, 11, 10, 18, 22, 16, 13, 14, 10, 26, 28, 14, 10, 10)):
    ws.column_dimensions[col].width = w
ws2 = wb.create_sheet('Lignes')
ws2.append(['Commande', 'Société', 'Client', 'Article', 'Famille', 'Qté cdée', 'Qté livrée', 'Qté facturée', 'Reste à facturer HT', 'Montant ligne HT'])
for o in orders:
    for l in byo[o['id']]:
        if l['invoice_status'] == 'to invoice':
            ws2.append([o['name'], o['company_id'][1], o['partner_id'][1], (l['product_id'] and l['product_id'][1]) or l['name'], famille(l),
                        l['product_uom_qty'], l['qty_delivered'], l['qty_invoiced'], round(l['untaxed_amount_to_invoice'], 2), round(l['price_subtotal'], 2)])
for cell in ws2[1]:
    cell.font = Font(bold=True)
ws2.freeze_panes = 'A2'; ws2.auto_filter.ref = ws2.dimensions
ws3 = wb.create_sheet('Synthèse'); ws3.append(['Catégorie', 'Commandes', 'Reste à facturer HT'])
syn = collections.defaultdict(lambda: [0, 0.0])
for r in rows:
    syn[r['cat']][0] += 1; syn[r['cat']][1] += r['reste']
for k, v in sorted(syn.items(), key=lambda kv: -kv[1][1]):
    ws3.append([k, v[0], round(v[1], 2)])
ws3.append(['TOTAL', len(rows), round(sum(r['reste'] for r in rows), 2)])
ws3.append([]); ws3.append(['Société', 'Commandes', 'Reste à facturer HT'])
sc = collections.defaultdict(lambda: [0, 0.0])
for r in rows:
    sc[r['soc']][0] += 1; sc[r['soc']][1] += r['reste']
for k, v in sorted(sc.items(), key=lambda kv: -kv[1][1]):
    ws3.append([k, v[0], round(v[1], 2)])
for col, w in zip('ABC', (60, 12, 20)):
    ws3.column_dimensions[col].width = w
out = 'C:/Users/xavfe/Desktop/Maquignon/Etat_commandes_a_facturer_%s.xlsx' % today.strftime('%Y-%m-%d')
wb.save(out)
print('fichier :', out)
print('synthèse :')
for k, v in sorted(syn.items(), key=lambda kv: -kv[1][1]):
    print('  %-70s %4d cdes  %12.2f EUR HT' % (k, v[0], v[1]))
print('  TOTAL reste à facturer : %.2f EUR HT' % sum(r['reste'] for r in rows))
print('par société :', {k: (v[0], round(v[1], 2)) for k, v in sc.items()})
old = [r for r in rows if r['age'] > 90]
print('> 90 jours :', len(old), 'commandes, %.2f EUR HT' % sum(r['reste'] for r in old))

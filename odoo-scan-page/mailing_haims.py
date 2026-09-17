# -*- coding: utf-8 -*-
"""Mailings « nouveaux tarifs au 01/10/2026 » de Carrière d'Haims : un mailing par liste de prix (Par défaut = particuliers,
Tarif Pro = travaux publics). Destinataires = clients facturés par Haims depuis le 01/01/2025, selon leur liste de prix.
dry   : listes + textes dans un Excel sur le Bureau
apply : idem + création des deux mailings en BROUILLON dans Odoo (Marketing par e-mail), PDF joint, rien n'est envoyé."""
import os, ssl, sys, re, io, base64, datetime, collections, xmlrpc.client
sys.stdout.reconfigure(encoding='utf-8')
mode = (sys.argv[1] if len(sys.argv) > 1 else 'dry').lower()
U = 'https://maquignon.odoo.com'; D = 'maquignon'; us = os.environ['ODOO_USER']; p = os.environ['ODOO_PWD']
c = ssl.create_default_context()
uid = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/common', context=c).authenticate(D, us, p, {})
m = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/object', context=c)
x = lambda mo, me, *a, **k: m.execute_kw(D, uid, p, mo, me, list(a), k)
ctx = {'allowed_company_ids': [4, 1, 2, 3, 13], 'active_test': False}   # Haims en société courante, accès aux fiches des autres sociétés
PAR_DEFAUT, TARIF_PRO = 271, 281
PDF = {'Particulier': 'C:/Users/xavfe/Downloads/Tarifs part Haims 2026 octobre.pdf', 'Travaux Publics': 'C:/Users/xavfe/Downloads/Tarifs pro Haims 2026 octobre.pdf'}
today = datetime.date.today()
# ---------- destinataires
inv = x('account.move', 'read_group', [['company_id', '=', 4], ['move_type', '=', 'out_invoice'], ['state', '=', 'posted'], ['invoice_date', '>=', '2025-01-01']],
        ['amount_untaxed:sum'], ['commercial_partner_id'], lazy=False, context=ctx)
fact = {r['commercial_partner_id'][0]: (r['__count'], r['amount_untaxed']) for r in inv}
last = {r['commercial_partner_id'][0]: (r.get('invoice_date:max') or r.get('invoice_date') or '') for r in x('account.move', 'read_group', [['company_id', '=', 4], ['move_type', '=', 'out_invoice'], ['state', '=', 'posted']], ['invoice_date:max'], ['commercial_partner_id'], lazy=False, context=ctx)}
parts = x('res.partner', 'read', sorted(fact), fields=['name', 'email', 'phone', 'city', 'specific_property_product_pricelist', 'is_company', 'x_studio_civilit_type_socit'], context=ctx)
ok = re.compile(r'^[^@\s]+@[^@\s]+\.[a-z]{2,}$', re.I)
groupes = {'Particulier': [], 'Travaux Publics': [], 'Autre liste': []}
sans = []
for pt in parts:
    pl = pt['specific_property_product_pricelist'] and pt['specific_property_product_pricelist'][0]
    g = 'Travaux Publics' if pl == TARIF_PRO else ('Particulier' if pl in (PAR_DEFAUT, None, False) else 'Autre liste')
    row = {'id': pt['id'], 'nom': pt['name'], 'email': (pt['email'] or '').strip(), 'tel': pt['phone'] or '', 'ville': pt['city'] or '', 'liste': (pt['specific_property_product_pricelist'] and pt['specific_property_product_pricelist'][1]) or 'Par défaut (EUR)',
           'n_fact': fact[pt['id']][0], 'ca': round(fact[pt['id']][1], 2), 'derniere': last.get(pt['id'], ''), 'groupe': g}
    (groupes[g] if ok.match(row['email']) else sans).append(row)
print('destinataires avec e-mail :', {k: len(v) for k, v in groupes.items()}, '| sans e-mail valide :', len(sans), dict(collections.Counter(r['groupe'] for r in sans)))
# ---------- textes
SIGN = """<p>Carrière d'Haims — Groupe Maquignon<br/>51 rue du Prieuré, 86230 Usseau<br/>Tél. 05 49 84 18 99 — carriere-d-haims@orange.fr</p>"""
TEXTES = {
    'Particulier': {
        'sujet': "Carrière d'Haims : nouveaux tarifs au 1er octobre 2026",
        'html': """<p>Bonjour,</p>
<p>Vous trouverez ci-joint notre nouvelle grille de tarifs particuliers, applicable au <b>1er octobre 2026</b>, pour l'ensemble de nos granulats au départ de la carrière d'Haims.</p>
<p>Comme l'ensemble de la profession, nous devons répercuter la hausse de nos coûts d'exploitation, en particulier l'énergie et le carburant nécessaires à l'extraction, au concassage et au chargement. Nous avons veillé à contenir cette évolution au strict nécessaire.</p>
<p>Les prix s'entendent hors taxes, à la tonne, au départ de la carrière ; la TGAP et l'éco-contribution s'ajoutent selon la réglementation en vigueur. Le transport reste sur demande, selon la destination et le tonnage.</p>
<p>Notre équipe reste à votre disposition pour tout devis ou conseil sur le choix des matériaux.</p>
<p>Nous vous remercions de votre confiance.</p>""" + SIGN},
    'Travaux Publics': {
        'sujet': "Carrière d'Haims : tarifs Travaux Publics applicables au 1er octobre 2026",
        'html': """<p>Madame, Monsieur,</p>
<p>Vous trouverez ci-joint notre grille de tarifs <b>Travaux Publics</b> applicable au <b>1er octobre 2026</b>, pour l'ensemble de nos granulats au départ de la carrière d'Haims.</p>
<p>Cette révision répercute la hausse de nos coûts de production, principalement l'énergie, le carburant et la maintenance des installations, dans un contexte que vous connaissez. Nous l'avons limitée au strict nécessaire afin de préserver la compétitivité de vos chantiers.</p>
<p>Les prix s'entendent hors taxes, à la tonne, au départ de la carrière, hors TGAP et éco-contribution. Les conditions particulières convenues avec certains d'entre vous sont mises à jour dans les mêmes proportions ; les transports restent chiffrés sur demande selon la destination et le tonnage.</p>
<p>Les commandes en cours et les devis signés avant le 1er octobre restent facturés aux conditions convenues.</p>
<p>Nous vous remercions de votre fidélité et restons à votre disposition.</p>""" + SIGN},
}
# ---------- Excel
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='01666B')
first = True
for g in ('Particulier', 'Travaux Publics', 'Autre liste'):
    if g == 'Autre liste' and not groupes[g]:
        continue
    ws = wb.active if first else wb.create_sheet(); first = False
    ws.title = 'Mailing ' + g
    ws.append(['Client', 'E-mail', 'Téléphone', 'Ville', 'Liste de prix', 'Factures depuis 2025', 'CA HT depuis 2025', 'Dernière facture'])
    for cell in ws[1]:
        cell.font = H; cell.fill = HF
    for r in sorted(groupes[g], key=lambda r: -r['ca']):
        ws.append([r['nom'], r['email'], r['tel'], r['ville'], r['liste'], r['n_fact'], r['ca'], r['derniere']])
    for col, w in zip('ABCDEFGH', (36, 36, 16, 18, 20, 10, 12, 12)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
ws = wb.create_sheet('Sans e-mail')
ws.append(['Client', 'Mailing', 'Téléphone', 'Ville', 'Liste de prix', 'Factures depuis 2025', 'CA HT depuis 2025', 'Dernière facture'])
for cell in ws[1]:
    cell.font = H; cell.fill = HF
for r in sorted(sans, key=lambda r: -r['ca']):
    ws.append([r['nom'], r['groupe'], r['tel'], r['ville'], r['liste'], r['n_fact'], r['ca'], r['derniere']])
for col, w in zip('ABCDEFGH', (36, 16, 16, 18, 20, 10, 12, 12)):
    ws.column_dimensions[col].width = w
ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
wt = wb.create_sheet('Textes'); wt.column_dimensions['A'].width = 18; wt.column_dimensions['B'].width = 130
for g, t in TEXTES.items():
    wt.append([g, 'Objet : ' + t['sujet']]); wt.cell(wt.max_row, 1).font = Font(bold=True)
    wt.append(['', re.sub(r'<[^>]+>', '', t['html'].replace('</p>', '\n').replace('<br/>', '\n')).strip()]); wt.cell(wt.max_row, 2).alignment = Alignment(wrap_text=True, vertical='top')
    wt.append(['', 'Pièce jointe : ' + os.path.basename(PDF[g])]); wt.append([])
out = 'C:/Users/xavfe/Desktop/Maquignon/Mailing_tarifs_Haims_2026-10_%s.xlsx' % today.strftime('%Y-%m-%d')
wb.save(out); print('fichier :', out)
if mode != 'apply':
    sys.exit(0)
# ---------- brouillons dans Odoo (Marketing par e-mail), modèle Contacts, domaine = les destinataires
model_id = x('ir.model', 'search', [['model', '=', 'res.partner']])[0]
soc = x('res.company', 'read', [4], fields=['name', 'email'])[0]
for g in ('Particulier', 'Travaux Publics'):
    ids = [r['id'] for r in groupes[g]]
    if not ids:
        continue
    pdf_b64 = base64.b64encode(io.open(PDF[g], 'rb').read()).decode()
    att = x('ir.attachment', 'create', [{'name': os.path.basename(PDF[g]), 'datas': pdf_b64, 'mimetype': 'application/pdf', 'res_model': 'mailing.mailing', 'res_id': 0}], context=ctx)
    att = att[0] if isinstance(att, list) else att
    vals = {'subject': TEXTES[g]['sujet'], 'body_html': TEXTES[g]['html'], 'body_arch': TEXTES[g]['html'], 'mailing_model_id': model_id,
            'mailing_domain': str([['id', 'in', ids]]), 'email_from': '%s <%s>' % (soc['name'], soc['email']), 'reply_to': soc['email'],
            'attachment_ids': [[6, 0, [att]]], 'name': 'Tarifs Haims 01/10/2026 — ' + g}
    try:
        mid = x('mailing.mailing', 'create', [vals], context=ctx)
    except Exception as e:  # noqa: BLE001
        vals.pop('name', None); vals.pop('body_arch', None)
        mid = x('mailing.mailing', 'create', [vals], context=ctx)
    mid = mid[0] if isinstance(mid, list) else mid
    x('ir.attachment', 'write', [att], {'res_id': mid}, context=ctx)
    mm = x('mailing.mailing', 'read', [mid], fields=['state', 'subject', 'mailing_domain'], context=ctx)[0]
    print('mailing créé (brouillon) :', mid, mm['state'], mm['subject'], '| destinataires', len(ids))

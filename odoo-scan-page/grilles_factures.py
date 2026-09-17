# -*- coding: utf-8 -*-
"""Deux grilles de tarifs (Particulier / Professionnel) proposées à partir des factures clients validées de SARL MAQUIGNON
(prix unitaires nets réellement facturés, par article et par type de client, un poids par client pour que les gros
comptes négociés ne fassent pas la grille). Lecture seule sur Odoo.
Usage : python grilles_factures.py [AAAA-MM-JJ]  (date de début, défaut 2026-01-01)"""
import os, ssl, sys, re, xmlrpc.client, collections, datetime, statistics
sys.stdout.reconfigure(encoding='utf-8')
DEPUIS = sys.argv[1] if len(sys.argv) > 1 else '2026-01-01'
U = 'https://maquignon.odoo.com'; D = 'maquignon'; us = os.environ['ODOO_USER']; p = os.environ['ODOO_PWD']
c = ssl.create_default_context()
uid = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/common', context=c).authenticate(D, us, p, {})
m = xmlrpc.client.ServerProxy(U + '/xmlrpc/2/object', context=c)
x = lambda mo, me, *a, **k: m.execute_kw(D, uid, p, mo, me, list(a), k)
ctx = {'allowed_company_ids': [1], 'active_test': False}
today = datetime.date.today(); today_s = today.isoformat()
PUBLIC, PRO = 302, 299
# ---------- articles + prix de liste actuels (évaluation locale des règles, pas de champ price en 19)
prods = x('product.product', 'search_read', [['sale_ok', '=', True], ['company_id', 'in', [1, False]]],
          fields=['default_code', 'display_name', 'lst_price', 'standard_price', 'categ_id', 'uom_id', 'product_tmpl_id', 'active'], limit=5000, context=ctx)
prod = {pr['id']: pr for pr in prods}
items = x('product.pricelist.item', 'search_read', [['pricelist_id', 'in', [PUBLIC, PRO]]],
          fields=['pricelist_id', 'applied_on', 'product_tmpl_id', 'product_id', 'categ_id', 'compute_price', 'fixed_price', 'percent_price', 'base',
                  'base_pricelist_id', 'price_discount', 'price_surcharge', 'price_round', 'min_quantity', 'date_start', 'date_end'], limit=20000, context=ctx)
byp = collections.defaultdict(list)
for it in items:
    byp[it['pricelist_id'][0]].append(it)
cats = {ct['id']: ct['parent_path'] for ct in x('product.category', 'search_read', [], fields=['parent_path'], context=ctx)}
ORDER = {'0_product_variant': 0, '1_product': 1, '2_product_category': 2, '3_global': 3}


def prix(pl_id, pr, depth=0):
    if depth > 5:
        return pr['lst_price']
    its = [it for it in byp[pl_id] if (not it['date_start'] or it['date_start'][:10] <= today_s) and (not it['date_end'] or it['date_end'][:10] >= today_s) and (it['min_quantity'] or 0) <= 1]
    for it in sorted(its, key=lambda it: (ORDER[it['applied_on']], -(it['min_quantity'] or 0), -(it['categ_id'][0] if it['categ_id'] else 0), -it['id'])):
        a = it['applied_on']
        if a == '0_product_variant' and (not it['product_id'] or it['product_id'][0] != pr['id']):
            continue
        if a == '1_product' and (not it['product_tmpl_id'] or it['product_tmpl_id'][0] != pr['product_tmpl_id'][0]):
            continue
        if a == '2_product_category' and not (it['categ_id'] and pr['categ_id'] and cats.get(pr['categ_id'][0], '').startswith(cats.get(it['categ_id'][0], '#'))):
            continue
        base = it.get('base') or 'list_price'
        bp = prix(it['base_pricelist_id'][0], pr, depth + 1) if (base == 'pricelist' and it.get('base_pricelist_id')) else (pr['standard_price'] if base == 'standard_price' else pr['lst_price'])
        if it['compute_price'] == 'fixed':
            return it['fixed_price'] or 0.0
        if it['compute_price'] == 'percentage':
            return bp * (1 - (it['percent_price'] or 0) / 100)
        v = bp * (1 - (it.get('price_discount') or 0) / 100)
        if it.get('price_round'):
            v = round(v / it['price_round']) * it['price_round']
        return v + (it.get('price_surcharge') or 0)
    return pr['lst_price']


# ---------- lignes de factures
aml = x('account.move.line', 'search_read', [['move_id.move_type', '=', 'out_invoice'], ['parent_state', '=', 'posted'], ['company_id', '=', 1], ['date', '>=', DEPUIS],
                                             ['display_type', '=', 'product'], ['product_id', '!=', False], ['quantity', '>', 0]],
        fields=['product_id', 'quantity', 'price_unit', 'discount', 'price_subtotal', 'product_uom_id', 'partner_id', 'date', 'move_name'], limit=200000, context=ctx)
print('lignes de factures validées depuis %s : %d' % (DEPUIS, len(aml)))
pids = sorted({l['partner_id'][0] for l in aml if l['partner_id']})
CIV = 'x_studio_civilit_type_socit'
pfields = ['name', 'is_company', 'vat', 'commercial_partner_id', CIV]
parts = {pt['id']: pt for pt in x('res.partner', 'read', pids, fields=pfields, context=ctx)} if pids else {}
cps = sorted({pt['commercial_partner_id'][0] for pt in parts.values()} - set(parts))
for pt in (x('res.partner', 'read', cps, fields=pfields, context=ctx) if cps else []):
    parts[pt['id']] = pt
PART_CIV = {'M.', 'Mme', 'M. & Mme', 'M', 'Mr', 'Mlle', 'M. et Mme', 'Monsieur', 'Madame', 'Melle'}


def groupe(pid):
    pt = parts.get(pid)
    if not pt:
        return 'Professionnel'
    cp = parts.get(pt['commercial_partner_id'][0], pt)
    civ = (cp.get(CIV) or '').strip()
    if civ:
        return 'Particulier' if civ in PART_CIV else 'Professionnel'
    return 'Professionnel' if (cp['is_company'] or cp['vat']) else 'Particulier'


obs = collections.defaultdict(list)   # (article, groupe) -> [(pu_net, qté, client, date, facture)]
uom_diff = 0; lignes = []
for l in aml:
    pr = prod.get(l['product_id'][0])
    if not pr:
        continue
    if l['product_uom_id'] and pr['uom_id'] and l['product_uom_id'][0] != pr['uom_id'][0]:
        uom_diff += 1; continue
    pu = round(l['price_subtotal'] / l['quantity'], 2)
    g = groupe(l['partner_id'][0]) if l['partner_id'] else 'Professionnel'
    cp = parts.get(l['partner_id'][0], {}).get('commercial_partner_id', [0, ''])[1] if l['partner_id'] else ''
    obs[(pr['id'], g)].append((pu, l['quantity'], cp, l['date'], l['move_name']))
    lignes.append((l['date'], l['move_name'], cp, g, pr['display_name'], l['quantity'], pu, round(l['price_subtotal'], 2)))
print('lignes hors unité de la fiche (ignorées) :', uom_diff, '| lignes par groupe :', dict(collections.Counter(r[3] for r in lignes)),
      '| clients par groupe :', {g: len({r[2] for r in lignes if r[3] == g}) for g in ('Particulier', 'Professionnel')})


def stats(o):
    """Statistiques par lignes ET par client (prix habituel de chaque client = son prix le plus fréquent)."""
    if not o:
        return None
    pus = sorted(v[0] for v in o); n = len(pus); q = sum(v[1] for v in o); ca = sum(v[0] * v[1] for v in o)
    mode, freq = collections.Counter(pus).most_common(1)[0]
    p25 = pus[max(0, int(0.25 * (n - 1)))]; p75 = pus[min(n - 1, int(0.75 * (n - 1)))]
    parcli = collections.defaultdict(list)
    for v in o:
        parcli[v[2]].append(v)
    cli = {}
    for k, vs in parcli.items():
        cli[k] = (collections.Counter(v[0] for v in vs).most_common(1)[0][0], len(vs), sum(v[1] for v in vs))
    cprix = sorted(v[0] for v in cli.values())
    cmode, cfreq = collections.Counter(cprix).most_common(1)[0]
    return {'n': n, 'qte': q, 'ca': ca, 'min': pus[0], 'p25': p25, 'med': statistics.median(pus), 'p75': p75, 'max': pus[-1], 'mode': mode, 'freq': freq,
            'moy': ca / q if q else statistics.median(pus), 'nc': len(cli), 'cmode': cmode, 'cfreq': cfreq, 'cmed': statistics.median(cprix),
            'top': sorted(cli.items(), key=lambda kv: -kv[1][1])[:3]}


def proposition(st, repli, libelle_repli):
    """Prix de grille : prix pratiqué par le plus grand nombre de clients (>= 2 clients et >= 30 % des clients), sinon médiane
    des prix par client ; avec 1 ou 2 clients seulement, médiane des lignes ; sans facture, prix actuel."""
    if not st:
        return repli, 'aucune facture : ' + libelle_repli
    if st['nc'] >= 3:
        if st['cfreq'] >= 2 and st['cfreq'] / st['nc'] >= 0.3:
            return st['cmode'], 'prix de %d clients sur %d' % (st['cfreq'], st['nc'])
        return round(st['cmed'], 2), 'médiane des prix par client (%d clients, %.2f à %.2f)' % (st['nc'], st['p25'], st['p75'])
    if st['n'] >= 3:
        return round(st['med'], 2), '%d client(s) seulement, médiane de %d lignes' % (st['nc'], st['n'])
    return repli, '%d ligne(s) seulement : %s' % (st['n'], libelle_repli)


# ---------- classeur
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
wb = Workbook()
H = Font(bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='01666B'); JAUNE = PatternFill('solid', fgColor='FFF2CC'); ROUGE = PatternFill('solid', fgColor='F8CBAD')


def entete(ws, cols, widths):
    ws.append(cols)
    for cell in ws[1]:
        cell.font = H; cell.fill = HF; cell.alignment = Alignment(wrap_text=True, vertical='center')
    ws.row_dimensions[1].height = 34
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'


ws = wb.active; ws.title = 'Méthode'; ws.column_dimensions['A'].width = 125
for line in [
    'GRILLES DE TARIFS PROPOSÉES — SARL MAQUIGNON — factures validées du %s au %s' % (DEPUIS, today_s),
    '',
    'Source : lignes des factures clients validées (hors avoirs). Prix unitaire net = montant HT de la ligne / quantité (remises incluses), dans l\'unité de la fiche.',
    'Type de client : civilité de la fiche (M., Mme = Particulier ; SARL, SAS, EARL… = Professionnel) ; sans civilité : société ou n° de TVA = Professionnel.',
    'Prix proposé : chaque client compte pour un (son prix habituel = son prix le plus fréquent), pour que les gros comptes négociés ne fassent pas la grille.',
    '   - au moins 3 clients : le prix partagé par le plus grand nombre de clients (au moins 2 clients et 30 % d\'entre eux), sinon la médiane des prix par client ;',
    '   - 1 ou 2 clients : médiane des lignes ; aucune facture : prix actuel conservé (fiche pour Particulier, Tarif Pro 2026 pour Professionnel).',
    '   - Particulier sans facture et sans prix sur la fiche : aligné sur le prix pro proposé, à relever si besoin.',
    'Colonne « Option : Particulier = Pro + 8 % » : rappel du rapport actuel entre tes listes (Tarif Pro 2026 = prix fiche - 7,4 % sur les pierres), si tu veux garder un écart fixe.',
    'Signaux : « prix dispersés » (P75 - P25 > 25 % de la médiane : prix au devis), « pro >= particulier » (à arbitrer), « particuliers loin de la fiche » (> 15 %).',
    'Unité facturée : les pré-sciées et tranches à épaisseur unique « (N cm) » sont facturées au m² alors que la fiche est en m³ : prix fiche et Tarif Pro sont ramenés au m² (prix/m³ × épaisseur) ; la grille est donc au m² pour ces variantes.',
    'Onglets Détail : par article, nombre de lignes et de clients, quantités, CA, min / quartiles / médiane / max, prix le plus fréquent, principaux clients et leur prix.',
    'Les colonnes jaunes PARTICULIER / PROFESSIONNEL retenus sont pré-remplies avec la proposition : corrige-les, c\'est ce que je chargerai dans Odoo.',
]:
    ws.append([line])
ws['A1'].font = Font(bold=True, size=13)
wg = wb.create_sheet('Grille proposée')
entete(wg, ['Réf.', 'Article', 'Catégorie', 'Unité fiche', 'Unité facturée', 'Prix fiche actuel (unité fiche)', 'Prix fiche ramené à l’unité facturée', 'Tarif Pro 2026 ramené à l’unité facturée', 'CA facturé période', 'Particuliers : clients / lignes', 'Particuliers : prix observé',
            'Pros : clients / lignes', 'Pros : prix observé', 'PARTICULIER retenu', 'PROFESSIONNEL retenu', 'Remise pro %', 'Option : Particulier = Pro + 8 %', 'Base de la proposition', 'Signaux', 'Pros : principaux clients (prix, lignes)'],
       (12, 48, 28, 8, 9, 11, 11, 11, 12, 11, 12, 11, 12, 13, 15, 9, 13, 52, 34, 60))
rows_detail = {'Particulier': [], 'Professionnel': []}
n_art = 0
for pr in sorted(prods, key=lambda pr: -sum(st['ca'] for st in (stats(obs.get((pr['id'], g))) for g in ('Particulier', 'Professionnel')) if st)):
    sp = stats(obs.get((pr['id'], 'Particulier'))); so = stats(obs.get((pr['id'], 'Professionnel')))
    if not sp and not so and not pr['active']:
        continue
    fiche0 = round(pr['lst_price'], 2); pro0 = round(prix(PRO, pr), 2)
    ep = re.search(r'\((\d+) cm\)', pr['display_name'])
    coef = int(ep.group(1)) / 100 if (ep and pr['uom_id'][1] == 'm³') else 1
    unite_fact = 'm²' if coef != 1 else pr['uom_id'][1]
    fiche = round(fiche0 * coef, 2); pro_act = round(pro0 * coef, 2); pro_act_aff = pro_act if abs(pro_act - fiche) > 0.005 else None
    pro_p, base_o = proposition(so, pro_act_aff, 'Tarif Pro 2026 conservé' if pro_act_aff else 'pas de prix pro spécifique')
    part_p, base_p = proposition(sp, fiche, 'prix fiche conservé')
    if (not sp or sp['n'] < 3) and fiche <= 1 and pro_p:
        part_p, base_p = pro_p, 'pas de facture particulier ni de prix fiche : aligné sur le prix pro, à relever'
    signaux = []
    for g, st in (('particuliers', sp), ('pros', so)):
        if st and st['n'] >= 3 and st['med'] and (st['p75'] - st['p25']) / st['med'] > 0.25:
            signaux.append('prix dispersés %s' % g)
    if part_p and pro_p and pro_p >= part_p - 0.005:
        signaux.append('pro >= particulier')
    if sp and fiche > 1 and abs(sp['cmed'] - fiche) / fiche > 0.15:
        signaux.append('particuliers loin de la fiche')
    if not sp and not so:
        signaux.append('aucune facture sur la période')
    ca = (sp['ca'] if sp else 0) + (so['ca'] if so else 0)
    wg.append([pr['default_code'] or '', pr['display_name'] + ('' if pr['active'] else ' (archivé)'), (pr['categ_id'] and pr['categ_id'][1]) or '', pr['uom_id'][1], unite_fact, fiche0, fiche, pro_act_aff,
               round(ca, 2), '%d / %d' % (sp['nc'], sp['n']) if sp else '', sp['cmode'] if sp else None, '%d / %d' % (so['nc'], so['n']) if so else '', so['cmode'] if so else None,
               part_p, pro_p, None, None, 'Part. : %s | Pro : %s' % (base_p, base_o), ', '.join(signaux),
               ' ; '.join('%s %.2f (%d l.)' % (k[:28], v[0], v[1]) for k, v in so['top']) if so else ''])
    r = wg.max_row; n_art += 1
    wg.cell(r, 14).fill = JAUNE; wg.cell(r, 15).fill = JAUNE
    wg.cell(r, 16).value = '=IF(AND(N%d>0,O%d<>""),ROUND((N%d-O%d)/N%d*100,1),"")' % (r, r, r, r, r)
    wg.cell(r, 17).value = '=IF(O%d<>"",ROUND(O%d*1.08,2),"")' % (r, r)
    if 'pro >= particulier' in signaux:
        wg.cell(r, 19).fill = ROUGE
    elif signaux:
        wg.cell(r, 19).fill = JAUNE
    for g, st in (('Particulier', sp), ('Professionnel', so)):
        if st:
            rows_detail[g].append([pr['default_code'] or '', pr['display_name'], unite_fact, st['n'], st['nc'], round(st['qte'], 2), round(st['ca'], 2), st['min'], st['p25'], st['med'], st['p75'], st['max'],
                                   st['mode'], '%d/%d' % (st['freq'], st['n']), st['cmode'], '%d/%d' % (st['cfreq'], st['nc']), round(st['moy'], 2), fiche, pro_act_aff,
                                   ' ; '.join('%s %.2f (%d l.)' % (k[:28], v[0], v[1]) for k, v in st['top'])])
wg.auto_filter.ref = wg.dimensions
for g in ('Particulier', 'Professionnel'):
    wd = wb.create_sheet('Détail ' + ('particuliers' if g == 'Particulier' else 'professionnels'))
    entete(wd, ['Réf.', 'Article', 'Unité', 'Lignes', 'Clients', 'Qté', 'CA HT', 'Min', 'P25', 'Médiane', 'P75', 'Max', 'Prix le + fréquent (lignes)', 'Fréq.', 'Prix le + fréquent (clients)', 'Fréq. clients', 'Moyenne pondérée', 'Prix fiche', 'Tarif Pro 2026', 'Principaux clients (prix, lignes)'],
           (12, 48, 8, 7, 7, 10, 12, 9, 9, 9, 9, 9, 11, 8, 11, 8, 10, 10, 10, 60))
    for row in sorted(rows_detail[g], key=lambda r: -r[6]):
        wd.append(row)
    wd.auto_filter.ref = wd.dimensions
wl = wb.create_sheet('Lignes')
entete(wl, ['Date', 'Facture', 'Client', 'Type', 'Article', 'Qté', 'PU net HT', 'Montant HT'], (11, 16, 34, 13, 48, 9, 10, 11))
for row in sorted(lignes, key=lambda r: (r[4], r[0])):
    wl.append(list(row))
wl.auto_filter.ref = wl.dimensions
out = 'C:/Users/xavfe/Desktop/Maquignon/Grilles_tarifs_proposees_Maquignon_%s.xlsx' % today.strftime('%Y-%m-%d')
wb.save(out)
print('fichier :', out, '| articles dans la grille :', n_art)
print('\n%-46s %8s %8s | %7s %9s | %8s %9s | %s' % ('article', 'fiche', 'pro act.', 'part c/l', 'PART prop', 'pro c/l', 'PRO prop', 'signaux'))
for row in list(wg.iter_rows(min_row=2, max_row=28, values_only=True)):
    print('%-46s %8s %8s | %7s %9s | %8s %9s | %s' % (row[1][:46], row[6], row[7] or '', row[9], row[13], row[11], row[14] or '', row[18]))

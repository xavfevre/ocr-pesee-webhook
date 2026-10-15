# -*- coding: utf-8 -*-
"""« Tarifs pierres » : import des tarifs pierres par fichier Excel, page web /tarifs-pierres (Xavier, 15/10/2026).

Modèle Excel généré depuis Odoo (prix actuels), rempli puis réimporté : pour chaque pierre, prix PRO et PARTICULIER en
deux phases (dates choisies sur la page, ex. 01/11/2026 puis 01/01/2027), plus une colonne par client à prix négociés
(en-tête = nom du client dans Odoo ; une colonne ajoutée = un nouveau client). Aperçu avant chargement, puis écriture :
  - règles à prix fixe datées sur les listes PRO et PARTICULIER cibles (paramètre `maquignon.tarifs_pierres_cfg`),
  - règles à prix fixe sans date sur la liste dédiée de chaque client (créée au besoin, repli = formule 100 % de sa liste
    de base, même mécanique que la page « Tarifs client »).
Unités : les listes Odoo sont au m³ ; les pierres d'épaisseur « (N cm) » se vendent au m² (le devis multiplie par
l'épaisseur) : les prix du fichier sont dans l'unité du devis et convertis en €/m³ (÷ épaisseur) à l'écriture.
Autres actions : création des listes cibles, bascule des clients des listes 2026 vers 2027, grille PDF par client.
Accès : ?k=<clé maquignon.tarifs_key> (même jeton que la page Tarifs client).
"""
import io
import json
import os
import re
import threading
import time
import xmlrpc.client
from datetime import date, datetime, timedelta, timezone

from flask import Blueprint, jsonify, request, send_file, send_from_directory

try:
    from zoneinfo import ZoneInfo
    TZ = ZoneInfo('Europe/Paris')
except Exception:  # noqa: BLE001
    TZ = timezone(timedelta(hours=1))
UTC = timezone.utc

bp = Blueprint('tarifs_pierres', __name__)
ICI = os.path.dirname(os.path.abspath(__file__))
SOC = 1
PARAM_CFG = 'maquignon.tarifs_pierres_cfg'      # {"pro": id, "part": id, "pro_prec": 299, "part_prec": 302}
PARAM_DEFAUT_PL = 'maquignon.tarifs_defaut_pl'  # listes de référence par société de la page Tarifs client
PARAM_DEDIEES = 'maquignon.tarifs_pl_dediees'
LISTES_PARTAGEES = {299, 302, 12}
NOM_PRO, NOM_PART = 'Tarif Pro 2027', 'Tarif Particulier 2027'
RE_EP = re.compile(r'\((\d+)\s*cm\)')
CTX = {'context': {'allowed_company_ids': [SOC], 'active_test': False}}

_CONN = {}
_LOCK = threading.Lock()
_CLE = {'val': '', 't': 0.0}


def _call(model, method, *params, **kw):
    with _LOCK:
        if 'uid' not in _CONN:
            url, db, user, pwd = os.environ.get('ODOO_URL'), os.environ.get('ODOO_DB'), os.environ.get('ODOO_USER'), os.environ.get('ODOO_PASSWORD')
            uid = xmlrpc.client.ServerProxy(f'{url}/xmlrpc/2/common').authenticate(db, user, pwd, {})
            if not uid:
                raise ValueError('authentification Odoo refusée')
            _CONN.update(uid=uid, models=xmlrpc.client.ServerProxy(f'{url}/xmlrpc/2/object'), db=db, pwd=pwd)
        try:
            return _CONN['models'].execute_kw(_CONN['db'], _CONN['uid'], _CONN['pwd'], model, method, list(params), kw)
        except xmlrpc.client.Fault as e:
            if 'cannot marshal None' in str(e):
                return None
            raise


def _cle_ok():
    now = time.time()
    if now - _CLE['t'] > 600 or not _CLE['val']:
        try:
            _CLE['val'] = _call('ir.config_parameter', 'get_param', 'maquignon.tarifs_key', '') or ''
        except Exception:  # noqa: BLE001
            pass
        _CLE['t'] = now
    k = request.args.get('k') or request.form.get('k') or (request.get_json(silent=True) or {}).get('k') or ''
    return bool(_CLE['val']) and k == _CLE['val']


one = lambda v: v[0] if isinstance(v, list) else v  # noqa: E731
nombre = lambda v: v if isinstance(v, (int, float)) and not isinstance(v, bool) else None  # noqa: E731
nom_net = lambda s: re.sub(r'^\[[^\]]+\] ', '', s or '').replace(' (archivé)', '').strip()  # noqa: E731


def _cfg():
    try:
        c = json.loads(_call('ir.config_parameter', 'get_param', PARAM_CFG, '') or '{}')
    except Exception:  # noqa: BLE001
        c = {}
    c.setdefault('pro_prec', 299); c.setdefault('part_prec', 302)
    return c


def _cfg_ecrire(c):
    _call('ir.config_parameter', 'set_param', PARAM_CFG, json.dumps(c))


def _paris_utc(jour, fin=False):
    """'2026-11-01' -> datetime UTC du début (ou de la fin) de ce jour à Paris, format Odoo."""
    d = datetime.strptime(jour, '%Y-%m-%d')
    d = d.replace(hour=23, minute=59, second=59) if fin else d
    return d.replace(tzinfo=TZ).astimezone(UTC).strftime('%Y-%m-%d %H:%M:%S')


def _phases(d1, d2):
    """Bornes UTC des deux phases : [d1, d2 - 1 s] puis [d2, ∞)."""
    fin1 = (datetime.strptime(d2, '%Y-%m-%d') - timedelta(days=1)).strftime('%Y-%m-%d')
    return {'p1': (_paris_utc(d1), _paris_utc(fin1, fin=True)), 'p2': (_paris_utc(d2), False)}


# ---------------------------------------------------------------- pierres et prix courants
def _pierres():
    """Pierres vendables de la société (catégories Pierres*, unité m³) avec coefficient d'unité du devis."""
    cats = _call('product.category', 'search', [['complete_name', 'ilike', 'Pierres']], **CTX)
    prods = _call('product.product', 'search_read', [['categ_id', 'in', cats], ['sale_ok', '=', True], ['active', '=', True], ['product_tmpl_id.company_id', 'in', [SOC, False]]],
                  fields=['default_code', 'display_name', 'lst_price', 'uom_name', 'categ_id', 'product_tmpl_id'], order='categ_id, default_code, id', **CTX)
    parents = {c['id']: c['parent_path'] for c in _call('product.category', 'read', cats, ['parent_path'], **CTX)}
    out = []
    for pr in prods:
        if (pr['uom_name'] or '') in ('Forfait', 'Unité', 'Units', 'Unit(s)'):
            continue                                   # articles à l'unité : sur devis
        mm = RE_EP.search(pr['display_name']) if pr['uom_name'] == 'm³' else None
        ep = int(mm.group(1)) / 100.0 if mm else None
        out.append({'id': pr['id'], 'code': pr['default_code'] or '', 'nom': nom_net(pr['display_name']), 'cat': pr['categ_id'][1].split('/')[-1].strip() if pr['categ_id'] else '',
                    'tmpl': pr['product_tmpl_id'] and pr['product_tmpl_id'][0], 'categs': tuple(int(i) for i in (parents.get(pr['categ_id'][0], '') or '').strip('/').split('/') if i) if pr['categ_id'] else (),
                    'fiche': pr['lst_price'], 'unite': 'm²' if ep else (pr['uom_name'] or ''), 'coef': (1.0 / ep) if ep else 1.0})   # prix m³ = prix unité × coef
    out.sort(key=lambda q: (q['cat'].lower(), q['code'] or '~', q['nom'].lower()))
    return out


_CACHE_REGLES = {}


def _regles(liste_id):
    """Règles d'une liste (toutes), mises en cache 60 s."""
    c = _CACHE_REGLES.get(liste_id)
    if c and time.time() - c[0] < 60:
        return c[1]
    r = _call('product.pricelist.item', 'search_read', [['pricelist_id', '=', liste_id]],
              fields=['applied_on', 'product_id', 'product_tmpl_id', 'categ_id', 'compute_price', 'fixed_price', 'percent_price', 'price_discount', 'base', 'base_pricelist_id', 'date_start', 'date_end', 'min_quantity'], **CTX)
    _CACHE_REGLES[liste_id] = (time.time(), r)
    return r


def _prix(liste_id, pierre, quand, profondeur=0):
    """Prix €/m³ d'une pierre sur une liste à une date UTC 'YYYY-MM-DD HH:MM:SS' (règles variante, puis formule générale)."""
    if not liste_id or profondeur > 4:
        return pierre['fiche']
    rs = _regles(liste_id)
    def ok(r):
        return (not r['date_start'] or r['date_start'] <= quand) and (not r['date_end'] or r['date_end'] >= quand) and (r['min_quantity'] or 0) <= 1
    cand = [r for r in rs if r['applied_on'] == '0_product_variant' and r['product_id'] and r['product_id'][0] == pierre['id'] and r['compute_price'] == 'fixed' and ok(r)]
    if cand:
        return sorted(cand, key=lambda r: -r['id'])[0]['fixed_price']
    glob = [r for r in rs if r['applied_on'] == '1_product' and r['product_tmpl_id'] and r['product_tmpl_id'][0] == pierre.get('tmpl') and ok(r)] or \
           [r for r in rs if r['applied_on'] == '2_product_category' and r.get('categ_id') and r['categ_id'][0] in pierre.get('categs', ()) and ok(r)] or \
           [r for r in rs if r['applied_on'] == '3_global' and ok(r)]
    if glob:
        g = sorted(glob, key=lambda r: -r['id'])[0]
        if g['compute_price'] == 'formula' and g['base'] == 'pricelist' and g['base_pricelist_id']:
            base = _prix(g['base_pricelist_id'][0], pierre, quand, profondeur + 1)
        else:
            base = pierre['fiche']
        if g['compute_price'] == 'percentage':
            return base * (1 - (g['percent_price'] or 0) / 100.0)
        if g['compute_price'] == 'formula':
            return base * (1 - (g['price_discount'] or 0) / 100.0)
        if g['compute_price'] == 'fixed':
            return g['fixed_price']
    return pierre['fiche']


# ---------------------------------------------------------------- clients à prix négociés
def _dediees():
    try:
        return set(json.loads(_call('ir.config_parameter', 'get_param', PARAM_DEDIEES, '') or '[]'))
    except Exception:  # noqa: BLE001
        return set()


def _simple(t):
    import unicodedata
    t = unicodedata.normalize('NFKD', t or '').encode('ascii', 'ignore').decode().lower()
    return re.sub(r'[^a-z0-9]', '', t)


def _liste_dediee(pl_id, pl_name, partner_name, n_utilisateurs=None):
    if not pl_id or pl_id in LISTES_PARTAGEES:
        return False
    if pl_id in _dediees() or n_utilisateurs == 1:
        return True
    a, b = _simple(partner_name), _simple(pl_name)
    return bool(a) and bool(b) and (a[:6] in b or b[:6] in a)


def _clients_negocies(pierres):
    """Clients ayant une liste dédiée avec au moins un prix fixe sur une pierre : {partner_id: {...}}."""
    pid_de = {p['id']: p for p in pierres}
    listes = _call('product.pricelist', 'search_read', [['company_id', 'in', [SOC, False]], ['id', 'not in', list(LISTES_PARTAGEES)]], fields=['name'], **CTX)
    out = {}
    for l in listes:
        parts = _call('res.partner', 'search_read', [['specific_property_product_pricelist', '=', l['id']], ['parent_id', '=', False]], fields=['name'], limit=3, **CTX)
        if not parts:
            continue
        pp = parts[0]
        if not _liste_dediee(l['id'], l['name'], pp['name'], len(parts)):
            continue
        lignes = {}
        for r in _regles(l['id']):
            if r['applied_on'] == '0_product_variant' and r['product_id'] and r['product_id'][0] in pid_de and r['compute_price'] == 'fixed' and not r['date_start']:
                q = pid_de[r['product_id'][0]]; lignes[q['id']] = round(r['fixed_price'] / q['coef'], 2)
        base = next((r['base_pricelist_id'][1] for r in _regles(l['id']) if r['applied_on'] == '3_global' and r['base'] == 'pricelist' and r['base_pricelist_id']), '')
        if lignes:
            out[pp['id']] = {'partner_id': pp['id'], 'nom': pp['name'], 'liste_id': l['id'], 'liste': l['name'], 'base': base, 'lignes': lignes}
    return out


# ---------------------------------------------------------------- modèle Excel
def _modele(d1, d2):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    pierres = _pierres(); cfg = _cfg(); clients = _clients_negocies(pierres)
    auj = datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')
    pro_l, part_l = cfg.get('pro') or cfg['pro_prec'], cfg.get('part') or cfg['part_prec']
    wb = Workbook(); ws = wb.active; ws.title = 'Tarifs'
    gras = Font(bold=True); gris = Font(color='666666', italic=True); bleu = PatternFill('solid', fgColor='EEF2FF'); jaune = PatternFill('solid', fgColor='FFF8DC'); mauve = PatternFill('solid', fgColor='F3E8FF')
    ws.append(['Tarifs pierres : modèle d import (SARL Maquignon)']); ws['A1'].font = Font(bold=True, size=13)
    ws.append(["Prix HT dans l'unité du devis (m² pour les pierres d'épaisseur « (N cm) », m³ sinon). Colonnes jaunes : prix à charger, phase 1 à partir du %s et phase 2 à partir du %s (dates réglables sur la page). "
               "Colonnes mauves : prix négociés des clients (sans date) ; pour ajouter un client, ajouter une colonne avec son nom exact Odoo avant « Remarque ». Case vide = pas de changement." % (d1, d2)])
    ws['A2'].font = gris; ws['A2'].alignment = Alignment(wrap_text=True, vertical='top'); ws.row_dimensions[2].height = 44
    ws.append([])
    noms_cli = [c['nom'] for c in sorted(clients.values(), key=lambda c: c['nom'].lower())]
    tete = ['Code', 'Pierre', 'Unité', 'PRO actuel', 'PRO phase 1', 'PRO phase 2', 'PARTICULIER actuel', 'PARTICULIER phase 1', 'PARTICULIER phase 2'] + noms_cli + ['Remarque']
    ws.append(tete)
    for i, c in enumerate(ws[4], start=1):
        c.font = gras; c.fill = mauve if 9 < i <= 9 + len(noms_cli) else bleu; c.alignment = Alignment(wrap_text=True, vertical='center', horizontal='center')
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(tete)); ws.row_dimensions[4].height = 36
    cli_par_nom = {c['nom']: c for c in clients.values()}
    cat = None
    for q in pierres:
        if q['cat'] != cat:
            cat = q['cat']; ws.append([cat]); ws.cell(ws.max_row, 1).font = gras; ws.merge_cells(start_row=ws.max_row, start_column=1, end_row=ws.max_row, end_column=len(tete))
        ligne = [q['code'], q['nom'], q['unite'], round(_prix(pro_l, q, auj) / q['coef'], 2), None, None, round(_prix(part_l, q, auj) / q['coef'], 2), None, None]
        ligne += [cli_par_nom[n]['lignes'].get(q['id']) for n in noms_cli] + ['']
        ws.append(ligne); r = ws.max_row
        for col in (5, 6, 8, 9):
            ws.cell(r, col).fill = jaune
        for i in range(len(noms_cli)):
            ws.cell(r, 10 + i).fill = mauve
        for col in range(4, 10 + len(noms_cli)):
            ws.cell(r, col).number_format = '# ##0.00'
    for i, wd in enumerate([13, 44, 6, 11, 11, 11, 12, 12, 12] + [12] * len(noms_cli) + [26]):
        ws.column_dimensions[ws.cell(4, i + 1).column_letter].width = wd
    ws.freeze_panes = 'D5'
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    return buf


# ---------------------------------------------------------------- lecture d'un fichier
def _lire(fichier):
    """Lit un fichier (modèle ou tableau de validation) : colonnes repérées par leur titre en ligne 4."""
    from openpyxl import load_workbook
    wb = load_workbook(fichier, data_only=True)
    ws = wb['Tarifs'] if 'Tarifs' in wb.sheetnames else (wb['Validation'] if 'Validation' in wb.sheetnames else wb.active)
    titres = [(c, str(ws.cell(4, c).value or '').strip()) for c in range(1, ws.max_column + 1)]
    def cols(prefixe):
        return [c for c, t in titres if t.upper().startswith(prefixe) and not any(m in t.lower() for m in ('actuel', 'liste', 'factur', 'hausse'))]
    c_pro, c_part = cols('PRO')[:2], cols('PARTICULIER')[:2]
    fixes = {'code', 'pierre', 'unité', 'unite', 'remarque', 'validé (oui / non)', 'validé'}
    c_cli = [(c, t) for c, t in titres if t and t.lower() not in fixes and c not in c_pro and c not in c_part and not t.upper().startswith(('PRO', 'PARTICULIER', 'HAUSSE'))]
    pierres = _pierres(); par_code = {q['code']: q for q in pierres if q['code']}; par_nom = {q['nom'].lower(): q for q in pierres}
    lignes, inconnues = [], []
    for r in range(5, ws.max_row + 1):
        code, nom, unite = (ws.cell(r, c).value for c in (1, 2, 3))
        if not nom or not unite:
            continue
        q = par_code.get(str(code).strip()) if code else None
        q = q or par_nom.get(nom_net(str(nom)).lower())
        if not q:
            inconnues.append('%s %s' % (code or '', nom)); continue
        val = lambda c: nombre(ws.cell(r, c).value) if c else None  # noqa: E731
        lignes.append({'id': q['id'], 'code': q['code'], 'nom': q['nom'], 'unite': q['unite'], 'coef': q['coef'],
                       'pro1': val(c_pro[0] if len(c_pro) > 0 else None), 'pro2': val(c_pro[1] if len(c_pro) > 1 else None),
                       'part1': val(c_part[0] if len(c_part) > 0 else None), 'part2': val(c_part[1] if len(c_part) > 1 else None),
                       'clients': {t: nombre(ws.cell(r, c).value) for c, t in c_cli if nombre(ws.cell(r, c).value)}})
    return {'lignes': lignes, 'inconnues': inconnues, 'colonnes_clients': [t for _, t in c_cli], 'entetes': {'pro': [t for c, t in titres if c in c_pro], 'part': [t for c, t in titres if c in c_part]}}


def _partenaire(nom):
    """Partenaire d'une colonne client : nom exact, sinon ilike unique."""
    F = ['name', 'specific_property_product_pricelist', 'is_company']
    c = _call('res.partner', 'search_read', [['name', '=ilike', nom], ['parent_id', '=', False]], fields=F, limit=5, **CTX)
    if len(c) > 1:      # homonymes : celui qui a une liste dédiée, sinon le plus facturé depuis 2025
        ded = [p for p in c if p['specific_property_product_pricelist'] and _liste_dediee(p['specific_property_product_pricelist'][0], p['specific_property_product_pricelist'][1], p['name'])]
        if len(ded) == 1:
            return ded[0]
        c = sorted(c, key=lambda p: -_call('account.move', 'search_count', [['partner_id', 'child_of', p['id']], ['move_type', '=', 'out_invoice'], ['state', '=', 'posted'], ['invoice_date', '>=', '2025-01-01']], **CTX))[:1]
    if not c:
        c = _call('res.partner', 'search_read', [['name', 'ilike', nom], ['parent_id', '=', False], ['is_company', '=', True]], fields=F, limit=3, **CTX)
    return c[0] if len(c) == 1 else (None if not c else {'ambigu': [p['name'] for p in c]})


def _apercu(fichier, d1, d2):
    lu = _lire(fichier); cfg = _cfg(); ph = _phases(d1, d2); auj = datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')
    pro_l, part_l = cfg.get('pro'), cfg.get('part')
    par_id = {q['id']: q for q in _pierres()}
    for l in lu['lignes']:
        q = par_id[l['id']]
        l['pro_act'] = round(_prix(pro_l or cfg['pro_prec'], q, auj) / l['coef'], 2); l['part_act'] = round(_prix(part_l or cfg['part_prec'], q, auj) / l['coef'], 2)
        for k in ('pro1', 'pro2', 'part1', 'part2'):
            l[k + '_m3'] = round(l[k] * l['coef'], 2) if l[k] else None
    clients = []
    for t in lu['colonnes_clients']:
        p = _partenaire(t); n = sum(1 for l in lu['lignes'] if l['clients'].get(t))
        if not p:
            clients.append({'colonne': t, 'etat': 'introuvable', 'n': n}); continue
        if 'ambigu' in p:
            clients.append({'colonne': t, 'etat': 'ambigu : ' + ', '.join(p['ambigu']), 'n': n}); continue
        pl = p['specific_property_product_pricelist']; pl_id = pl and pl[0]
        n_util = pl_id and _call('res.partner', 'search_count', [['specific_property_product_pricelist', '=', pl_id], ['parent_id', '=', False]], **CTX)
        dediee = pl and _liste_dediee(pl_id, pl[1], p['name'], n_util)
        clients.append({'colonne': t, 'partner_id': p['id'], 'nom': p['name'], 'liste': pl and pl[1] or 'aucune', 'liste_id': pl_id, 'etat': 'liste dédiée existante' if dediee else 'liste dédiée à créer', 'dediee': bool(dediee),
                        'base': 'part' if pl_id == (part_l or cfg['part_prec']) or pl_id == cfg['part_prec'] else 'pro', 'n': n})
    n_regles = sum(1 for l in lu['lignes'] for k in ('pro1', 'pro2', 'part1', 'part2') if l[k])
    return {'phases': {'d1': d1, 'd2': d2, 'utc': ph}, 'lignes': lu['lignes'], 'inconnues': lu['inconnues'], 'clients': clients, 'entetes': lu['entetes'],
            'listes': {'pro': pro_l, 'part': part_l}, 'n_regles': n_regles, 'n_clients': sum(c['n'] for c in clients if c.get('partner_id'))}


# ---------------------------------------------------------------- écriture
def _regle_fixe(liste_id, product_id, prix_m3, debut, fin):
    """Crée ou met à jour la règle variante à prix fixe de la liste pour ce produit et ce début de validité."""
    dom = [['pricelist_id', '=', liste_id], ['product_id', '=', product_id], ['applied_on', '=', '0_product_variant'], ['compute_price', '=', 'fixed'], ['date_start', '=', debut or False]]
    ex = _call('product.pricelist.item', 'search', dom, **CTX)
    vals = {'fixed_price': prix_m3, 'date_end': fin or False}
    if ex:
        _call('product.pricelist.item', 'write', ex, vals, **CTX); return 'maj'
    vals.update(pricelist_id=liste_id, product_id=product_id, applied_on='0_product_variant', compute_price='fixed', date_start=debut or False)
    _call('product.pricelist.item', 'create', [vals], **CTX); return 'creee'


def _liste_creer(nom, base_liste):
    devise = _call('product.pricelist', 'read', [base_liste], ['currency_id'], **CTX)[0]['currency_id'][0]
    lid = one(_call('product.pricelist', 'create', [{'name': nom, 'company_id': SOC, 'currency_id': devise}], **CTX))
    _call('product.pricelist.item', 'create', [{'pricelist_id': lid, 'applied_on': '3_global', 'compute_price': 'formula', 'base': 'pricelist', 'base_pricelist_id': base_liste, 'price_discount': 0}], **CTX)
    return lid


def _charger(fichier, d1, d2):
    ap = _apercu(fichier, d1, d2); cfg = _cfg(); ph = ap['phases']['utc']
    if not cfg.get('pro') or not cfg.get('part'):
        return {'erreur': 'Les listes cibles PRO et PARTICULIER ne sont pas définies : créer les listes 2027 d abord.'}
    stats = {'regles': 0, 'maj': 0, 'clients': 0, 'listes_creees': [], 'regles_clients': 0}
    for l in ap['lignes']:
        for lid, k1, k2 in ((cfg['pro'], 'pro1', 'pro2'), (cfg['part'], 'part1', 'part2')):
            for k, (deb, fin) in ((k1, ph['p1']), (k2, ph['p2'])):
                if l[k]:
                    res = _regle_fixe(lid, l['id'], l[k + '_m3'], deb, fin); stats['regles' if res == 'creee' else 'maj'] += 1
    for c in ap['clients']:
        if not c.get('partner_id') or not c['n']:
            continue
        lid = c['liste_id'] if c['dediee'] else None
        if not lid:
            base = cfg['part'] if c['base'] == 'part' else cfg['pro']
            lid = _liste_creer('Tarif %s' % c['nom'][:60], base); stats['listes_creees'].append(c['nom'])
            _call('res.partner', 'write', [c['partner_id']], {'specific_property_product_pricelist': lid}, **CTX)
            ded = _dediees(); ded.add(lid); _call('ir.config_parameter', 'set_param', PARAM_DEDIEES, json.dumps(sorted(ded)))
        for l in ap['lignes']:
            v = l['clients'].get(c['colonne'])
            if v:
                _regle_fixe(lid, l['id'], round(v * l['coef'], 2), False, False); stats['regles_clients'] += 1
        stats['clients'] += 1
    _CACHE_REGLES.clear()
    return {'ok': True, 'stats': stats, 'phases': ap['phases']}


# ---------------------------------------------------------------- grille d'un client (données communes PDF / HTML / e-mail)
def _pierres_tarifees(cfg, ph):
    """Pierres dont un prix de phase existe sur la liste PRO ou PARTICULIER cible (= la grille à diffuser)."""
    ids = set()
    for lid in (cfg.get('pro'), cfg.get('part')):
        for r in (_regles(lid) if lid else []):
            if r['applied_on'] == '0_product_variant' and r['product_id'] and r['date_start'] in (ph['p1'][0], ph['p2'][0]):
                ids.add(r['product_id'][0])
    return [q for q in _pierres() if q['id'] in ids]


def _grille_client(pid, d1, d2):
    pp = _call('res.partner', 'read', [pid], ['name', 'email', 'specific_property_product_pricelist'], **CTX)[0]
    cfg = _cfg(); ph = _phases(d1, d2); pl = pp['specific_property_product_pricelist']; pl_id = pl and pl[0]
    n_util = pl_id and _call('res.partner', 'search_count', [['specific_property_product_pricelist', '=', pl_id], ['parent_id', '=', False]], **CTX)
    dediee = bool(pl_id and _liste_dediee(pl_id, pl[1], pp['name'], n_util))
    liste = pl_id or cfg.get('pro') or cfg['pro_prec']
    nego = {r['product_id'][0]: r['fixed_price'] for r in _regles(liste) if dediee and r['applied_on'] == '0_product_variant' and r['product_id'] and r['compute_price'] == 'fixed' and not r['date_start']} if dediee else {}
    lignes = []
    for q in _pierres_tarifees(cfg, ph):
        pn = nego.get(q['id'])
        v1, v2 = (pn, pn) if pn else (_prix(liste, q, ph['p1'][0]), _prix(liste, q, ph['p2'][0]))
        lignes.append({'cat': q['cat'], 'nom': q['nom'], 'unite': q['unite'], 'v1': round(v1 / q['coef'], 2), 'v2': round(v2 / q['coef'], 2), 'nego': round(pn / q['coef'], 2) if pn else None})
    return {'partner': pp, 'lignes': lignes, 'nego': bool(nego), 'd1': d1, 'd2': d2, 'liste': pl and pl[1] or ''}


_fmt_prix = lambda v: ('%.2f' % v).replace('.', ',').replace(',00', '')  # noqa: E731
_jour = lambda s_: datetime.strptime(s_, '%Y-%m-%d').strftime('%d/%m/%Y')  # noqa: E731


def _grille_html(g):
    """Tableau HTML (styles en ligne, compatible e-mail) de la grille d'un client."""
    fin1 = (datetime.strptime(g['d2'], '%Y-%m-%d') - timedelta(days=1)).strftime('%d/%m/%Y')
    th = 'style="padding:6px 8px;border-bottom:2px solid #cbd5e1;font-size:12px;color:#475569;text-align:right;background:#f8fafc"'
    td = 'style="padding:5px 8px;border-bottom:1px solid #e2e8f0;font-size:13px;text-align:right;white-space:nowrap"'
    tdg = td.replace('text-align:right;white-space:nowrap', 'text-align:left')
    h = ['<table cellpadding="0" cellspacing="0" style="border-collapse:collapse;width:100%;font-family:Arial,Helvetica,sans-serif;color:#0f172a">',
         '<tr><th %s>Pierre</th><th %s>Unité</th><th %s>Du %s au %s</th><th %s>À partir du %s</th>%s</tr>' % (th.replace('right', 'left'), th, th, _jour(g['d1']), fin1, th, _jour(g['d2']), ('<th %s>Votre prix négocié</th>' % th) if g['nego'] else '')]
    cat = None
    for l in g['lignes']:
        if l['cat'] != cat:
            cat = l['cat']; h.append('<tr><td colspan="%d" style="padding:7px 8px 3px;font-weight:bold;font-size:13px;color:#1d4ed8;border-bottom:1px solid #e2e8f0">%s</td></tr>' % (5 if g['nego'] else 4, cat))
        fond = ' background:#fff4e0;' if l['nego'] else ''
        h.append('<tr><td %s>%s</td><td %s>%s</td><td %s>%s €</td><td %s>%s €</td>%s</tr>' % (tdg[:-1] + fond + '"', l['nom'], td[:-1] + fond + '"', l['unite'], td[:-1] + fond + '"', _fmt_prix(l['v1']), td[:-1] + fond + '"', _fmt_prix(l['v2']),
                                                                              ('<td %s><b>%s €</b></td>' % (td[:-1] + fond + '"', _fmt_prix(l['nego'])) if l['nego'] else '<td %s></td>' % td) if g['nego'] else ''))
    h.append('</table>')
    return ''.join(h)


def _pdf_client(g, soc):
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    INK = colors.HexColor('#0f172a'); GREY = colors.HexColor('#475569'); LIGHT = colors.HexColor('#eef2ff'); LINE = colors.HexColor('#cbd5e1'); NEGO = colors.HexColor('#fff4e0')
    P = ParagraphStyle('p', fontName='Helvetica', fontSize=9.6, leading=13, textColor=INK); H1 = ParagraphStyle('h1', parent=P, fontName='Helvetica-Bold', fontSize=15, leading=19)
    SMALL = ParagraphStyle('s', parent=P, fontSize=8, leading=10.5, textColor=GREY); CELL = ParagraphStyle('c', parent=P, fontSize=8.8, leading=11); CELLB = ParagraphStyle('cb', parent=CELL, fontName='Helvetica-Bold')
    fin1 = (datetime.strptime(g['d2'], '%Y-%m-%d') - timedelta(days=1)).strftime('%d/%m/%Y')
    buf = io.BytesIO(); doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=14 * mm, bottomMargin=14 * mm, title='Tarifs pierres', author=soc['name'])
    s_ = [Paragraph(soc['name'], SMALL), Paragraph('Tarifs pierres : %s' % g['partner']['name'], H1),
          Paragraph('Prix hors taxes, départ carrière, en deux étapes : du <b>%s</b> au <b>%s</b>, puis à partir du <b>%s</b>. Unités : m² pour les pierres d\'épaisseur, m³ pour les massifs, tranches et blocs. Édité le %s.'
                    % (_jour(g['d1']), fin1, _jour(g['d2']), date.today().strftime('%d/%m/%Y')), P), Spacer(1, 4)]
    if g['nego']:
        s_ += [Paragraph('Les lignes sur fond orangé sont <b>vos prix négociés</b>, maintenus sur les deux périodes.', P), Spacer(1, 4)]
    tete = ['Pierre', 'Unité', 'Du %s' % _jour(g['d1']), 'À partir du %s' % _jour(g['d2'])] + (['Votre prix négocié'] if g['nego'] else [])
    lignes = [[Paragraph(t, CELLB) for t in tete]]; st = [('BACKGROUND', (0, 0), (-1, 0), LIGHT), ('LINEBELOW', (0, 0), (-1, -1), 0.3, LINE), ('TOPPADDING', (0, 0), (-1, -1), 2.5), ('BOTTOMPADDING', (0, 0), (-1, -1), 2.5), ('ALIGN', (2, 1), (-1, -1), 'RIGHT')]
    cat = None
    for l in g['lignes']:
        if l['cat'] != cat:
            cat = l['cat']; lignes.append([Paragraph('<b>%s</b>' % cat, CELL)] + [''] * (len(tete) - 1)); st += [('SPAN', (0, len(lignes) - 1), (-1, len(lignes) - 1)), ('BACKGROUND', (0, len(lignes) - 1), (-1, len(lignes) - 1), colors.HexColor('#f8fafc'))]
        row = [Paragraph(l['nom'], CELL), Paragraph(l['unite'], CELL), Paragraph(_fmt_prix(l['v1']), CELL), Paragraph(_fmt_prix(l['v2']), CELL)]
        if g['nego']:
            row.append(Paragraph('<b>%s</b>' % _fmt_prix(l['nego']) if l['nego'] else '', CELL))
            if l['nego']:
                st.append(('BACKGROUND', (0, len(lignes)), (-1, len(lignes)), NEGO))
        lignes.append(row)
    larg = [86 * mm, 12 * mm, 38 * mm, 38 * mm] if not g['nego'] else [72 * mm, 12 * mm, 33 * mm, 33 * mm, 30 * mm]
    t = Table(lignes, colWidths=larg, repeatRows=1); t.setStyle(TableStyle(st)); s_ += [t, Spacer(1, 6), Paragraph('Les autres articles ne changent pas. Votre contact : %s, %s.' % (soc['phone'] or '', soc['email'] or ''), SMALL)]
    doc.build(s_); buf.seek(0)
    return buf


# ---------------------------------------------------------------- mailing E-mail Marketing
CHAMPS_MAILING = {'x_tarif_token': ('Jeton grille tarifs', 'char'), 'x_tarif_html': ('Grille tarifs (HTML)', 'html')}
PARAM_MAILING = 'maquignon.tarifs_pierres_mailing'
SUJET_MAILING = 'Carrières Maquignon : vos tarifs pierres 2027'
RELAIS_URL = os.environ.get('RELAIS_URL', 'https://ocr-pesee-webhook.onrender.com')


def _champs_mailing():
    """Crée les champs personnalisés sur res.partner s'ils manquent."""
    mid = _call('ir.model', 'search', [['model', '=', 'res.partner']], **CTX)[0]
    crees = []
    for nom, (lib, ttype) in CHAMPS_MAILING.items():
        if not _call('ir.model.fields', 'search', [['model', '=', 'res.partner'], ['name', '=', nom]], **CTX):
            _call('ir.model.fields', 'create', [{'name': nom, 'model_id': mid, 'field_description': lib, 'ttype': ttype, 'state': 'manual'}], **CTX); crees.append(nom)
    return crees


def _cibles_mailing():
    """Clients pro acheteurs de pierre (2025-2026, société 1) avec e-mail, hors liste particulier."""
    cfg = _cfg(); cats = _call('product.category', 'search', [['complete_name', 'ilike', 'Pierres']], **CTX)
    g = _call('account.move.line', 'read_group', [['move_id.move_type', '=', 'out_invoice'], ['move_id.state', '=', 'posted'], ['move_id.invoice_date', '>=', '2025-01-01'], ['company_id', '=', SOC], ['product_id.categ_id', 'in', cats]], ['price_subtotal:sum'], ['partner_id'], lazy=False, **CTX)
    pids = [r['partner_id'][0] for r in g if r['partner_id']]
    comm = {q['id']: q['commercial_partner_id'][0] for q in _call('res.partner', 'read', pids, ['commercial_partner_id'], **CTX)} if pids else {}
    ca = {}
    for r in g:
        if r['partner_id']:
            ca[comm[r['partner_id'][0]]] = ca.get(comm[r['partner_id'][0]], 0) + r['price_subtotal']
    out = []
    for q in _call('res.partner', 'read', list(ca), ['name', 'email', 'specific_property_product_pricelist'], **CTX) if ca else []:
        l = q['specific_property_product_pricelist'] and q['specific_property_product_pricelist'][0]
        if l and l not in (cfg['part_prec'], cfg.get('part')) and q['email']:
            out.append({'id': q['id'], 'nom': q['name'], 'email': q['email'], 'ca': round(ca[q['id']])})
    return sorted(out, key=lambda c: -c['ca'])


def _corps_mailing(d1, d2, modele=None):
    """Corps du mailing (QWeb : <t t-out>, t-attf-href) : {'body_arch': …, 'body_html': …}. Mise en page du brouillon
    d'Isabelle (« Nouveaux tarifs 2026 ») si fourni, texte remplacé ; sinon corps simple."""
    fin1 = (datetime.strptime(d2, '%Y-%m-%d') - timedelta(days=1)).strftime('%d/%m/%Y')
    P = 'style="font-family:Arial,Helvetica,sans-serif;color:#000000;font-size:16px;margin:0 0 16px"'
    bloc = ('<p %s>Vous trouverez ci-dessous votre grille de tarifs pierres, propre à votre compte. Elle s\'applique en deux étapes : du <strong>%s au %s</strong>, puis à partir du <strong>%s</strong>. '
            'Les prix s\'entendent hors taxes, départ carrière, au m² pour les pierres d\'épaisseur et au m³ pour les massifs, tranches et blocs. Les autres articles ne changent pas. '
            'Les devis signés avant le %s restent honorés aux conditions convenues.</p>'
            '<t t-out="object.x_tarif_html"/>'
            '<p style="text-align:center;margin:18px 0"><a t-attf-href="%s/tarifs-pierres/grille/#{object.x_tarif_token}" style="background:#1d4ed8;color:#ffffff;padding:10px 18px;border-radius:8px;text-decoration:none;font-weight:bold;font-family:Arial,Helvetica,sans-serif">Télécharger votre grille en PDF</a></p>'
            % (P, _jour(d1), fin1, _jour(d2), _jour(d1), RELAIS_URL))

    def adapter(html):
        h = html.replace('Nouveaux tarifs 2026', 'Nouveaux tarifs pierres 2027')
        h = re.sub(r'À partir du 1er(?:&nbsp;| )octobre(?:&nbsp;| )2026', 'Du 1er novembre 2026, puis du 1er janvier 2027', h)
        h = h.replace('<t t-out="object.contact_address_complete">client</t>', '<t t-out="object.name">client</t>')
        h, n1 = re.subn(r'<p[^>]*>(?:(?!</p>).)*?augmentation de 6(?:(?!</p>).)*?</p>', lambda _m: bloc, h, count=1, flags=re.S)
        h = re.sub(r'<p[^>]*>(?:(?!</p>).)*?Tous les devis signés et validés avant cette date(?:(?!</p>).)*?</p>', '', h, count=1, flags=re.S)
        return h if n1 else None

    if modele:
        arch, html = adapter(modele.get('body_arch') or ''), adapter(modele.get('body_html') or '')
        if arch and html:
            return {'body_arch': arch, 'body_html': html}
    simple = ('<div style="font-family:Arial,Helvetica,sans-serif;color:#0f172a;max-width:680px;margin:0 auto"><h2 style="color:#1d4ed8">Nouveaux tarifs pierres 2027</h2><p %s>Madame, Monsieur,</p>'
              '<p %s>Nous vous remercions pour la confiance que vous nous accordez pour l\'approvisionnement de vos chantiers en pierre de tuffeau et autres pierres calcaires. '
              'Pour faire face à la hausse de nos coûts d\'exploitation et continuer à investir dans nos gisements, nous ajustons nos tarifs pierres en deux étapes.</p>%s'
              '<p %s>Nous restons à votre disposition pour vos devis et vos chantiers.</p><p %s>Bien cordialement,<br/>SARL Maquignon</p></div>' % (P, P, bloc, P, P))
    return {'body_arch': simple, 'body_html': simple}


def _preparer_mailing(d1, d2):
    import secrets
    crees = _champs_mailing(); cibles = _cibles_mailing(); cfg = _cfg()
    if not cfg.get('pro'):
        return {'erreur': 'listes cibles absentes : créer les listes 2027 et charger la grille d abord'}
    n_html = 0
    for c in cibles:
        pp = _call('res.partner', 'read', [c['id']], ['x_tarif_token'], **CTX)[0]
        vals = {'x_tarif_html': _grille_html(_grille_client(c['id'], d1, d2))}
        if not pp.get('x_tarif_token'):
            vals['x_tarif_token'] = secrets.token_urlsafe(15)
        _call('res.partner', 'write', [c['id']], vals, **CTX); n_html += 1
    # brouillon : copie du brouillon « Nouveaux tarifs 2026 » d'Isabelle (mise en page) si présent, sinon corps simple
    modele = _call('mailing.mailing', 'search_read', [['subject', '=', 'Nouveaux tarifs 2026'], ['state', '=', 'draft']], fields=['id', 'body_arch', 'body_html', 'email_from', 'reply_to'], limit=1, **CTX)
    corps = _corps_mailing(d1, d2, modele[0] if modele else None)
    modele_partner = _call('ir.model', 'search', [['model', '=', 'res.partner']], **CTX)[0]
    vals = {'subject': SUJET_MAILING, 'body_arch': corps['body_arch'], 'body_html': corps['body_html'], 'mailing_model_id': modele_partner, 'mailing_domain': json.dumps([['id', 'in', [c['id'] for c in cibles]]]),
            'email_from': (modele[0]['email_from'] if modele else '') or '"SARL MAQUIGNON" <contact@maquignon.com>', 'reply_to_mode': 'new', 'reply_to': (modele[0]['reply_to'] if modele else '') or 'contact@maquignon.com'}
    existant = _call('mailing.mailing', 'search', [['subject', '=', SUJET_MAILING], ['state', '=', 'draft']], **CTX)
    if existant:
        _call('mailing.mailing', 'write', existant[:1], vals, **CTX); mid = existant[0]; action = 'brouillon mis à jour'
    else:
        mid = one(_call('mailing.mailing', 'create', [vals], **CTX)); action = 'brouillon créé'
    _call('ir.config_parameter', 'set_param', PARAM_MAILING, json.dumps({'mailing': mid, 'd1': d1, 'd2': d2, 'n': len(cibles), 'quand': datetime.utcnow().isoformat()}))
    return {'ok': True, 'mailing': mid, 'action': action, 'cibles': len(cibles), 'grilles': n_html, 'champs_crees': crees, 'modele_isabelle': bool(modele),
            'url': '%s/odoo/action-mass_mailing.mailing_mailing_action_mail/%d' % (os.environ.get('ODOO_URL', ''), mid)}


# ---------------------------------------------------------------- routes
@bp.route('/grille/<jeton>')
def grille_publique(jeton):
    """Grille PDF d'un client, lien du mailing (jeton propre au client, sans clé)."""
    if not re.fullmatch(r'[A-Za-z0-9_-]{12,40}', jeton or ''):
        return 'lien invalide', 404
    try:
        pp = _call('res.partner', 'search_read', [['x_tarif_token', '=', jeton]], fields=['name'], limit=1, **CTX)
    except Exception:  # noqa: BLE001  (champ absent tant qu'aucun mailing n'a été préparé)
        pp = []
    if not pp:
        return 'lien inconnu ou expiré', 404
    try:
        m_ = json.loads(_call('ir.config_parameter', 'get_param', PARAM_MAILING, '') or '{}')
    except Exception:  # noqa: BLE001
        m_ = {}
    g = _grille_client(pp[0]['id'], m_.get('d1', '2026-11-01'), m_.get('d2', '2027-01-01'))
    soc = _call('res.company', 'read', [SOC], ['name', 'phone', 'email'], **CTX)[0]
    return send_file(_pdf_client(g, soc), as_attachment=False, download_name='Tarifs_pierres_%s.pdf' % re.sub(r'[^A-Za-z0-9]+', '_', pp[0]['name'])[:50], mimetype='application/pdf')


@bp.route('/api/mailing/etat')
def mailing_etat():
    if not _cle_ok():
        return jsonify({'erreur': 'clé'}), 403
    cibles = _cibles_mailing()
    try:
        m_ = json.loads(_call('ir.config_parameter', 'get_param', PARAM_MAILING, '') or '{}')
    except Exception:  # noqa: BLE001
        m_ = {}
    etat_m = _call('mailing.mailing', 'read', [m_['mailing']], ['state', 'sent', 'subject'], **CTX)[0] if m_.get('mailing') else None
    return jsonify({'cibles': len(cibles), 'ca': sum(c['ca'] for c in cibles), 'exemples': cibles[:5], 'mailing': m_, 'etat_mailing': etat_m,
                    'modele_isabelle': bool(_call('mailing.mailing', 'search', [['subject', '=', 'Nouveaux tarifs 2026'], ['state', '=', 'draft']], **CTX)),
                    'url': ('%s/odoo/action-mass_mailing.mailing_mailing_action_mail/%d' % (os.environ.get('ODOO_URL', ''), m_['mailing'])) if m_.get('mailing') else ''})


@bp.route('/api/mailing/preparer', methods=['POST'])
def mailing_preparer():
    if not _cle_ok():
        return jsonify({'erreur': 'clé'}), 403
    j = request.get_json(silent=True) or {}
    try:
        return jsonify(_preparer_mailing(j.get('d1', '2026-11-01'), j.get('d2', '2027-01-01')))
    except Exception as e:  # noqa: BLE001
        return jsonify({'erreur': 'préparation interrompue : %s' % e}), 500


@bp.route('/api/mailing/apercu')
def mailing_apercu():
    """Grille HTML d'un client telle qu'elle sera insérée dans l'e-mail (contrôle)."""
    if not _cle_ok():
        return 'clé', 403
    g = _grille_client(int(request.args.get('partner', 0)), request.args.get('d1', '2026-11-01'), request.args.get('d2', '2027-01-01'))
    return '<html><body style="font-family:Arial;max-width:680px;margin:20px auto"><h3>%s</h3>%s</body></html>' % (g['partner']['name'], _grille_html(g))


@bp.route('/')
def page():
    return send_from_directory(ICI, 'tarifs_pierres_page.html')


@bp.route('/api/etat')
def etat():
    if not _cle_ok():
        return jsonify({'erreur': 'clé'}), 403
    cfg = _cfg(); pierres = _pierres()
    listes = {l['id']: l['name'] for l in _call('product.pricelist', 'read', [i for i in (cfg.get('pro'), cfg.get('part'), cfg['pro_prec'], cfg['part_prec']) if i], ['name'], **CTX)}
    n_prec = {k: _call('res.partner', 'search_count', [['specific_property_product_pricelist', '=', cfg[k]]], **CTX) for k in ('pro_prec', 'part_prec')}
    n_cibles = {k: _call('res.partner', 'search_count', [['specific_property_product_pricelist', '=', cfg[k]]], **CTX) if cfg.get(k) else 0 for k in ('pro', 'part')}
    clients = _clients_negocies(pierres)
    return jsonify({'cfg': cfg, 'listes': listes, 'n_pierres': len(pierres), 'clients_prec': n_prec, 'clients_cibles': n_cibles,
                    'clients_negocies': sorted(({**c, 'lignes': len(c['lignes'])} for c in clients.values()), key=lambda c: c['nom'].lower()),
                    'url_tarifs_client': '/tarifs-client/?k=%s&partner=' % _CLE['val']})


@bp.route('/api/modele.xlsx')
def modele():
    if not _cle_ok():
        return 'clé', 403
    d1, d2 = request.args.get('d1', '2026-11-01'), request.args.get('d2', '2027-01-01')
    return send_file(_modele(d1, d2), as_attachment=True, download_name='Tarifs_pierres_modele_%s.xlsx' % date.today().isoformat(), mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@bp.route('/api/apercu', methods=['POST'])
def apercu():
    if not _cle_ok():
        return jsonify({'erreur': 'clé'}), 403
    f = request.files.get('fichier')
    if not f:
        return jsonify({'erreur': 'fichier manquant'}), 400
    try:
        return jsonify(_apercu(io.BytesIO(f.read()), request.form.get('d1', '2026-11-01'), request.form.get('d2', '2027-01-01')))
    except Exception as e:  # noqa: BLE001
        return jsonify({'erreur': 'lecture impossible : %s' % e}), 400


@bp.route('/api/charger', methods=['POST'])
def charger():
    if not _cle_ok():
        return jsonify({'erreur': 'clé'}), 403
    f = request.files.get('fichier')
    if not f:
        return jsonify({'erreur': 'fichier manquant'}), 400
    try:
        return jsonify(_charger(io.BytesIO(f.read()), request.form.get('d1', '2026-11-01'), request.form.get('d2', '2027-01-01')))
    except Exception as e:  # noqa: BLE001
        _CACHE_REGLES.clear()
        return jsonify({'erreur': 'chargement interrompu : %s' % e}), 500


@bp.route('/api/listes', methods=['POST'])
def listes_creer():
    """Crée les listes cibles 2027 si absentes (repli = formule 100 % de la liste 2026 / du prix fiche) et les enregistre."""
    if not _cle_ok():
        return jsonify({'erreur': 'clé'}), 403
    cfg = _cfg(); faits = []
    for k, nom, prec in (('pro', NOM_PRO, cfg['pro_prec']), ('part', NOM_PART, cfg['part_prec'])):
        if cfg.get(k):
            continue
        ex = _call('product.pricelist', 'search', [['name', '=', nom], ['company_id', 'in', [SOC, False]]], **CTX)
        if ex:
            cfg[k] = ex[0]; faits.append('%s existait' % nom); continue
        devise = _call('product.pricelist', 'read', [prec], ['currency_id'], **CTX)[0]['currency_id'][0]
        lid = one(_call('product.pricelist', 'create', [{'name': nom, 'company_id': SOC, 'currency_id': devise, 'sequence': 5 if k == 'pro' else 6}], **CTX))
        regle = {'pricelist_id': lid, 'applied_on': '3_global', 'compute_price': 'formula', 'price_discount': 0}
        regle.update({'base': 'pricelist', 'base_pricelist_id': prec} if k == 'pro' else {'base': 'list_price'})
        _call('product.pricelist.item', 'create', [regle], **CTX); cfg[k] = lid; faits.append('%s créée (%d)' % (nom, lid))
    _cfg_ecrire(cfg); _CACHE_REGLES.clear()
    return jsonify({'ok': True, 'cfg': cfg, 'faits': faits})


@bp.route('/api/bascule', methods=['POST'])
def bascule():
    """Déplace les clients des listes 2026 vers les listes 2027, liste par défaut et liste de référence de la page Tarifs client."""
    if not _cle_ok():
        return jsonify({'erreur': 'clé'}), 403
    cfg = _cfg()
    if not cfg.get('pro') or not cfg.get('part'):
        return jsonify({'erreur': 'listes cibles absentes'}), 400
    n = {}
    for k in ('pro', 'part'):
        ids = _call('res.partner', 'search', [['specific_property_product_pricelist', '=', cfg[k + '_prec']]], **CTX)
        for i in range(0, len(ids), 200):
            _call('res.partner', 'write', ids[i:i + 200], {'specific_property_product_pricelist': cfg[k]}, **CTX)
        n[k] = len(ids)
    champ = _call('ir.model.fields', 'search', [['model', '=', 'res.partner'], ['name', '=', 'property_product_pricelist']], **CTX)
    defaut = _call('ir.default', 'search', [['field_id', 'in', champ], ['company_id', 'in', [SOC, False]]], **CTX) if champ else []
    if defaut:
        _call('ir.default', 'write', defaut, {'json_value': str(cfg['part'])}, **CTX)
    try:
        ref = json.loads(_call('ir.config_parameter', 'get_param', PARAM_DEFAUT_PL, '') or '{}')
    except Exception:  # noqa: BLE001
        ref = {}
    ref[str(SOC)] = cfg['pro']; _call('ir.config_parameter', 'set_param', PARAM_DEFAUT_PL, json.dumps(ref))
    # listes dédiées dont le repli est encore l'ancienne liste pro (12) ou Tarif Pro 2026 (299) -> repli sur Tarif Pro 2027
    rebasees = []
    for r in _call('product.pricelist.item', 'search_read', [['applied_on', '=', '3_global'], ['compute_price', '=', 'formula'], ['base', '=', 'pricelist'], ['base_pricelist_id', 'in', [12, cfg['pro_prec']]], ['pricelist_id', 'not in', [cfg['pro'], cfg['part']]]], fields=['pricelist_id'], **CTX):
        _call('product.pricelist.item', 'write', [r['id']], {'base_pricelist_id': cfg['pro']}, **CTX); rebasees.append(r['pricelist_id'][1])
    _CACHE_REGLES.clear()
    return jsonify({'ok': True, 'deplaces': n, 'defaut_maj': bool(defaut), 'rebasees': rebasees})


@bp.route('/api/pdf')
def pdf():
    """Grille PDF personnalisée d'un client (prix des deux phases sur sa liste, lignes négociées en orangé)."""
    if not _cle_ok():
        return 'clé', 403
    g = _grille_client(int(request.args.get('partner', 0)), request.args.get('d1', '2026-11-01'), request.args.get('d2', '2027-01-01'))
    soc = _call('res.company', 'read', [SOC], ['name', 'phone', 'email'], **CTX)[0]
    return send_file(_pdf_client(g, soc), as_attachment=True, download_name='Tarifs_pierres_%s.pdf' % re.sub(r'[^A-Za-z0-9]+', '_', g['partner']['name'])[:50], mimetype='application/pdf')

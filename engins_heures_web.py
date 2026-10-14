# -*- coding: utf-8 -*-
"""Heures de fonctionnement des engins non routiers : page web /heures-engins et ses services (Xavier, 14/10/2026).

Source : relevés d'horamètre du Parc automobile Odoo (fleet.vehicle.odometer ; pour les engins l'unité « miles » tient
lieu d'heures). Engins = catégories Pelle, Chargeuse, Tombereau, Chariot élévateur, ou compteur en « miles ».
La page affiche, pour une période (trimestre / semestre / année), le dernier relevé de chaque mois par engin, le
relevé de référence avant la période, le relevé de fin et la différence = heures de la période ; export Excel ;
saisie d'un nouveau relevé (écrit fleet.vehicle.odometer). Accès : ?k=<clé maquignon.engins_key>.
"""
import io
import os
import re
import threading
import time
import xmlrpc.client
from datetime import date, datetime, timedelta

from flask import Blueprint, jsonify, request, send_file, send_from_directory

bp = Blueprint('engins_heures', __name__)
ICI = os.path.dirname(os.path.abspath(__file__))
CATEGORIES = ('Pelle', 'Chargeuse', 'Tombereau', 'Chariot élévateur')
MOIS = ['janvier', 'février', 'mars', 'avril', 'mai', 'juin', 'juillet', 'août', 'septembre', 'octobre', 'novembre', 'décembre']
CTX = {'context': {'allowed_company_ids': [1, 3, 4], 'active_test': False}}

_CONN = {}
_LOCK = threading.Lock()
_CLE = {'val': '', 't': 0.0}


def _call(model, method, *params, **kw):
    with _LOCK:
        if 'uid' not in _CONN:
            url, db, user, pwd = os.environ.get('ODOO_URL'), os.environ.get('ODOO_DB'), os.environ.get('ODOO_USER'), os.environ.get('ODOO_PASSWORD')
            uid = xmlrpc.client.ServerProxy(f'{url}/xmlrpc/2/common').authenticate(db, user, pwd, {})
            if not uid:
                raise ValueError('authentification Odoo refusée')
            _CONN.update(uid=uid, models=xmlrpc.client.ServerProxy(f'{url}/xmlrpc/2/object'), db=db, pwd=pwd)
        try:
            return _CONN['models'].execute_kw(_CONN['db'], _CONN['uid'], _CONN['pwd'], model, method, list(params), kw)
        except xmlrpc.client.Fault as e:
            if 'cannot marshal None' in str(e):
                return None
            raise


def _cle_ok():
    now = time.time()
    if now - _CLE['t'] > 600 or not _CLE['val']:
        try:
            _CLE['val'] = _call('ir.config_parameter', 'get_param', 'maquignon.engins_key', '') or ''
        except Exception:  # noqa: BLE001
            pass
        _CLE['t'] = now
    k = request.args.get('k') or (request.get_json(silent=True) or {}).get('k') or ''
    return bool(_CLE['val']) and k == _CLE['val']


def _periode(periodicite, annee, index):
    """(debut, fin, libellé, [(annee, mois), ...]) de la période demandée."""
    if periodicite == 'annee':
        debut, fin, lib = date(annee, 1, 1), date(annee, 12, 31), 'Année %d' % annee
    elif periodicite == 'semestre':
        index = 1 if index < 1 else 2 if index > 2 else index
        debut, fin = (date(annee, 1, 1), date(annee, 6, 30)) if index == 1 else (date(annee, 7, 1), date(annee, 12, 31))
        lib = '%der semestre %d' % (index, annee) if index == 1 else '2e semestre %d' % annee
    else:
        periodicite = 'trimestre'; index = min(max(index, 1), 4)
        debut = date(annee, 3 * (index - 1) + 1, 1); fin = date(annee, 3 * index, [31, 30, 30, 31][index - 1])
        lib = '%der trimestre %d' % (index, annee) if index == 1 else '%de trimestre %d' % (index, annee)
    mois = []
    m = debut
    while m <= fin:
        mois.append((m.year, m.month)); m = date(m.year + (m.month == 12), m.month % 12 + 1, 1)
    return debut, fin, lib, mois


def _engins(archives=False):
    rows = _call('fleet.vehicle', 'search_read', [], fields=['name', 'license_plate', 'category_id', 'odometer', 'odometer_unit', 'company_id', 'active', 'model_id'], limit=1000, **CTX)
    eng = [v for v in rows if (v['odometer_unit'] == 'miles' or (v['category_id'] and v['category_id'][1] in CATEGORIES)) and (archives or v['active'])]
    eng.sort(key=lambda v: ((v['company_id'] and v['company_id'][1]) or '', (v['category_id'] and v['category_id'][1]) or 'zzz', v['name']))
    return eng


def _releves(ids, jusqua):
    rows = _call('fleet.vehicle.odometer', 'search_read', [['vehicle_id', 'in', ids], ['date', '<=', jusqua.strftime('%Y-%m-%d')]],
                 fields=['vehicle_id', 'date', 'value'], limit=20000, order='date, id', **CTX)
    par = {}
    for r in rows:
        par.setdefault(r['vehicle_id'][0], []).append((datetime.strptime(r['date'], '%Y-%m-%d').date(), float(r['value']), r['id']))
    return par


def _fmt(d):
    return d.strftime('%d/%m/%Y')


def _tableau(periodicite, annee, index, archives):
    debut, fin, lib, mois = _periode(periodicite, annee, index)
    eng = _engins(archives)
    rel = _releves([v['id'] for v in eng], fin)
    lignes, totaux = [], {}
    for v in eng:
        pts = rel.get(v['id'], [])
        par_mois = {}
        for d, val, rid in pts:
            if debut <= d <= fin:
                par_mois[(d.year, d.month)] = (d, val)          # trié par date : le dernier du mois reste
        avant = [p for p in pts if p[0] < debut]
        dans = [p for p in pts if debut <= p[0] <= fin]
        rem = []
        ref = avant[-1] if avant else None
        if ref is None and dans:
            ref = dans[0]; rem.append('pas de relevé avant la période : départ = premier relevé de la période (%s)' % _fmt(ref[0]))
        fin_rel = dans[-1] if dans else None
        if not dans:
            rem.append('aucun relevé pendant la période' + (' (dernier connu : %s)' % _fmt(avant[-1][0]) if avant else ''))
        elif fin_rel[0] < date(mois[-1][0], mois[-1][1], 1):
            rem.append('pas de relevé le dernier mois')
        diff, intervalle = None, None
        if ref and fin_rel and fin_rel[0] > ref[0]:
            diff = round(fin_rel[1] - ref[1]); jours = (fin_rel[0] - ref[0]).days
            intervalle = {'du': _fmt(ref[0]), 'au': _fmt(fin_rel[0]), 'jours': jours, 'exact': ref[0] >= debut - timedelta(days=45) and fin_rel[0] >= date(mois[-1][0], mois[-1][1], 1)}
            if diff < 0:
                rem.append('compteur en baisse : remise à zéro ou erreur de saisie, à corriger dans Odoo')
            elif diff > 24 * (jours or 1):
                rem.append('plus de 24 h par jour : erreur de saisie probable, à vérifier')
        if v['odometer_unit'] != 'miles':
            rem.append("unité « %s » dans Odoo (un engin doit être en « miles » = heures)" % v['odometer_unit'])
        soc = (v['company_id'] and v['company_id'][1]) or ''
        if diff is not None and diff >= 0:
            totaux[soc] = totaux.get(soc, 0) + diff
        lignes.append({'id': v['id'], 'nom': v['name'], 'categorie': (v['category_id'] and v['category_id'][1]) or '', 'societe': soc, 'actif': v['active'],
                       'mois': [{'date': _fmt(par_mois[m][0]), 'valeur': round(par_mois[m][1])} if m in par_mois else None for m in mois],
                       'ref': {'date': _fmt(ref[0]), 'valeur': round(ref[1])} if ref else None,
                       'fin': {'date': _fmt(fin_rel[0]), 'valeur': round(fin_rel[1])} if fin_rel else None,
                       'diff': diff, 'intervalle': intervalle, 'remarques': rem, 'dernier': {'date': _fmt(pts[-1][0]), 'valeur': round(pts[-1][1])} if pts else None})
    return {'periode': {'periodicite': periodicite, 'annee': annee, 'index': index, 'libelle': lib, 'debut': _fmt(debut), 'fin': _fmt(fin)},
            'mois': ['%s %d' % (MOIS[m - 1], a) for a, m in mois], 'lignes': lignes, 'totaux': totaux, 'total': sum(totaux.values())}


def _params():
    periodicite = request.args.get('periodicite') or 'trimestre'
    try:
        annee = int(request.args.get('annee') or date.today().year)
        index = int(request.args.get('index') or ((date.today().month - 1) // 3 + 1))
    except ValueError:
        annee, index = date.today().year, 1
    return periodicite, annee, index, request.args.get('archives') in ('1', 'true', 'oui')


@bp.route('/')
def page():
    return send_from_directory(ICI, 'engins_heures_page.html')


@bp.route('/api/tableau')
def tableau():
    if not _cle_ok():
        return jsonify({'error': 'clé invalide'}), 401
    return jsonify(_tableau(*_params()))


@bp.route('/api/releve', methods=['POST'])
def releve():
    """Nouveau relevé d'horamètre : {veh, date 'AAAA-MM-JJ', valeur}. Refusé s'il est inférieur au dernier relevé."""
    if not _cle_ok():
        return jsonify({'error': 'clé invalide'}), 401
    d = request.get_json(silent=True) or {}
    try:
        vid = int(d.get('veh')); val = float(str(d.get('valeur')).replace(',', '.')); jour = datetime.strptime(d.get('date') or date.today().isoformat(), '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return jsonify({'error': 'relevé incomplet (engin, date, heures)'}), 400
    if val < 0 or val > 200000:
        return jsonify({'error': 'valeur invraisemblable'}), 400
    if jour > date.today():
        return jsonify({'error': 'date dans le futur'}), 400
    dern = _call('fleet.vehicle.odometer', 'search_read', [['vehicle_id', '=', vid], ['date', '<=', jour.strftime('%Y-%m-%d')]], fields=['date', 'value'], order='date desc, id desc', limit=1, **CTX)
    if dern and val < dern[0]['value'] and not d.get('forcer'):
        return jsonify({'error': 'inférieur au dernier relevé (%s h le %s). Vérifiez la valeur ; pour un compteur remis à zéro, cochez « forcer ».' % (round(dern[0]['value']), _fmt(datetime.strptime(dern[0]['date'], '%Y-%m-%d'))), 'confirmable': True}), 409
    rid = _call('fleet.vehicle.odometer', 'create', [{'vehicle_id': vid, 'date': jour.strftime('%Y-%m-%d'), 'value': val}], **CTX)
    return jsonify({'ok': True, 'id': rid[0] if isinstance(rid, list) else rid})


@bp.route('/api/export.xlsx')
def export():
    if not _cle_ok():
        return jsonify({'error': 'clé invalide'}), 401
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    t = _tableau(*_params())
    wb = Workbook(); ws = wb.active; ws.title = re.sub(r'[\\/*?:\[\]]', ' ', t['periode']['libelle'])[:31]
    gras = Font(bold=True); fond = PatternFill('solid', fgColor='EEF2FF')
    ws['A1'] = 'Heures de fonctionnement des engins non routiers : %s (du %s au %s)' % (t['periode']['libelle'], t['periode']['debut'], t['periode']['fin']); ws['A1'].font = Font(bold=True, size=13)
    ws['A2'] = "Dernier relevé d'horamètre de chaque mois (Parc automobile Odoo) ; différence = relevé de fin - relevé de référence avant la période. Généré le %s." % _fmt(date.today())
    ent = ['Société', 'Catégorie', 'Engin', 'Relevé de référence', 'Date réf.'] + [m.capitalize() for m in t['mois']] + ['Relevé de fin', 'Date fin', 'Heures de la période', 'Jours entre les deux relevés', 'Remarques']
    ws.append([]); ws.append(ent)
    for c in ws[4]:
        c.font = gras; c.fill = fond; c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    for l in t['lignes']:
        ws.append([l['societe'], l['categorie'], l['nom'], l['ref']['valeur'] if l['ref'] else '', l['ref']['date'] if l['ref'] else '']
                  + [m['valeur'] if m else '' for m in l['mois']]
                  + [l['fin']['valeur'] if l['fin'] else '', l['fin']['date'] if l['fin'] else '', l['diff'] if l['diff'] is not None else '', l['intervalle']['jours'] if l['intervalle'] else '', ' ; '.join(l['remarques'])])
    ws.append([])
    col_h = 8 + len(t['mois'])
    for soc, h in t['totaux'].items():
        ws.append([soc, '', 'TOTAL ' + soc] + [''] * (col_h - 4) + [h]); ws.cell(ws.max_row, 3).font = gras; ws.cell(ws.max_row, col_h).font = gras
    ws.append(['', '', 'TOTAL engins'] + [''] * (col_h - 4) + [t['total']]); ws.cell(ws.max_row, 3).font = gras; ws.cell(ws.max_row, col_h).font = gras
    larg = [20, 16, 44, 12, 12] + [12] * len(t['mois']) + [12, 12, 14, 12, 70]
    for i, w in enumerate(larg, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row[3:len(larg) - 1]:
            c.number_format = '# ##0'
        row[-1].alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'D5'
    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    nom = 'Heures_engins_%s.xlsx' % re.sub(r'[^A-Za-z0-9]+', '_', t['periode']['libelle']).strip('_')
    return send_file(buf, as_attachment=True, download_name=nom, mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
